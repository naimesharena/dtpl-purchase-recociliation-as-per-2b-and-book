from __future__ import annotations

import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from gstr2b_recon import ReconciliationConfig, ReconciliationService, Transaction
from gstr2b_recon.storage import (
    load_outstanding_state,
    save_outstanding_state,
    write_excel_report,
)


class CarryForwardStateTests(unittest.TestCase):
    def test_open_item_round_trips_and_matches_a_later_period(self) -> None:
        config = ReconciliationConfig()
        service = ReconciliationService(config)
        book = Transaction(
            transaction_id="books:file:sheet:11",
            source="books",
            supplier_gstin="24ABCDE1234F1Z5",
            document_number="CN-55",
            document_type="credit_note",
            document_date="2025-09-15",
            books_month="2025-09",
            document_month="2025-09",
            taxable_value=Decimal("1000.00"),
            cgst=Decimal("90.00"),
            sgst=Decimal("90.00"),
            total_value=Decimal("1180.00"),
            classification="eligible",
            itc_eligible=True,
        )
        initial = service.reconcile([book], [], as_of_month="2025-09")
        self.assertEqual(len(initial.outstanding_items), 1)

        with tempfile.TemporaryDirectory() as directory:
            report = Path(directory) / "report.xlsx"
            write_excel_report(initial, report, config)
            from openpyxl import load_workbook

            workbook = load_workbook(report, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, [
                "Read Me", "Monthly Summary", "Reconciliation", "Books Not in 2B", "2B Not in Books", "Outstanding",
            ])
            self.assertEqual(workbook["Reconciliation"].max_row, 2)
            self.assertEqual(workbook["Outstanding"].max_row, 2)
            workbook.close()

            state = Path(directory) / "open_items.json"
            save_outstanding_state(state, initial.outstanding_items)
            loaded = load_outstanding_state(state)
            self.assertEqual(loaded[0].id, book.transaction_id)
            self.assertEqual(loaded[0].expected_gstr2b_month, "2025-10")
            self.assertEqual(loaded[0].normalized_cgst, Decimal("-90.00"))

            gstr = Transaction(
                transaction_id="2b:oct:note:7",
                source="gstr2b",
                supplier_gstin=book.supplier_gstin,
                document_number="CN-55",
                document_type="C",
                document_date="2025-09-15",
                document_month="2025-09",
                gstr2b_month="2025-10",
                taxable_value=Decimal("-1000.00"),
                cgst=Decimal("-90.00"),
                sgst=Decimal("-90.00"),
                total_value=Decimal("-1180.00"),
                itc_eligible=True,
            )
            later = service.reconcile([], [gstr], outstanding=loaded, as_of_month="2025-10")
            self.assertEqual(later.records[0].status, "CREDIT_NOTE_TIMING_DIFFERENCE")
            self.assertTrue(later.records[0].resolved)
            self.assertEqual(later.records[0].gstr2b.transaction_id, gstr.transaction_id)
            self.assertEqual(later.outstanding_items, [])


if __name__ == "__main__":
    unittest.main()
