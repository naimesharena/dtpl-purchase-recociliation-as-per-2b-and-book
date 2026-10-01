"""JSON carry-forward state and CSV/JSON report writers."""
from __future__ import annotations

import csv
import json
import os
import tempfile
from collections.abc import Iterable
from dataclasses import fields
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from .config import ReconciliationConfig
from .models import (
    OutstandingItem,
    ReconciliationRecord,
    ReconciliationResult,
    Transaction,
)
from .normalization import normalized_values

_STATE_VERSION = 1


def load_outstanding_state(path: str | Path | None) -> list[OutstandingItem]:
    if not path:
        return []
    state_path = Path(path)
    if not state_path.exists():
        return []
    with state_path.open(encoding="utf-8") as stream:
        data = json.load(stream)
    if data.get("version") != _STATE_VERSION:
        raise ValueError(f"Unsupported carry-forward state version in {state_path}")
    output: list[OutstandingItem] = []
    for item in data.get("outstanding_items", []):
        transaction = transaction_from_dict(item["transaction"])
        output.append(OutstandingItem(
            transaction=transaction,
            expected_gstr2b_month=item.get("expected_gstr2b_month"),
            normalized_taxable_value=_decimal(item.get("normalized_taxable_value", "0")),
            normalized_cgst=_decimal(item.get("normalized_cgst", "0")),
            normalized_sgst=_decimal(item.get("normalized_sgst", "0")),
            normalized_igst=_decimal(item.get("normalized_igst", "0")),
            normalized_cess=_decimal(item.get("normalized_cess", "0")),
            status=item.get("status", "TIMING_DIFFERENCE"),
            matched_transaction_id=item.get("matched_transaction_id"),
        ))
    return output


def save_outstanding_state(path: str | Path, items: Iterable[OutstandingItem]) -> None:
    state_path = Path(path)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": _STATE_VERSION,
        "outstanding_items": [outstanding_to_dict(item) for item in items],
    }
    _atomic_json_write(state_path, payload)


def outstanding_to_dict(item: OutstandingItem) -> dict[str, Any]:
    return {
        "id": item.id,
        "supplier_gstin": item.supplier_gstin,
        "document_number": item.document_number,
        "document_type": item.document_type,
        "document_date": item.document_date,
        "books_month": item.books_month,
        "gstr2b_month": item.gstr2b_month,
        "taxable_value": str(item.taxable_value),
        "cgst": str(item.cgst),
        "sgst": str(item.sgst),
        "igst": str(item.igst),
        "normalized_taxable_value": str(item.normalized_taxable_value),
        "normalized_cgst": str(item.normalized_cgst),
        "normalized_sgst": str(item.normalized_sgst),
        "normalized_igst": str(item.normalized_igst),
        "normalized_cess": str(item.normalized_cess),
        "status": item.status,
        "expected_gstr2b_month": item.expected_gstr2b_month,
        "matched_transaction_id": item.matched_transaction_id,
        "transaction": transaction_to_dict(item.transaction),
    }


def transaction_to_dict(transaction: Transaction) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for item in fields(Transaction):
        value = getattr(transaction, item.name)
        if isinstance(value, Decimal):
            value = str(value)
        result[item.name] = value
    return result


def transaction_from_dict(value: dict[str, Any]) -> Transaction:
    amounts = {"taxable_value", "cgst", "sgst", "igst", "cess", "total_value"}
    kwargs = dict(value)
    for key in amounts:
        kwargs[key] = _decimal(kwargs.get(key, "0"))
    return Transaction(**kwargs)


def write_reports(
    result: ReconciliationResult,
    output_directory: str | Path,
    config: ReconciliationConfig,
) -> dict[str, Path]:
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    paths = {
        "detail": output / "reconciliation_detail.csv",
        "section_a": output / "books_not_in_gstr2b.csv",
        "section_b": output / "gstr2b_not_in_books.csv",
        "summary": output / "monthly_summary.csv",
        "json": output / "reconciliation.json",
        "open_items": output / "outstanding_items.json",
        "workbook": output / "GSTR2B_vs_Books_Reconciliation.xlsx",
    }
    detail_rows = [_record_to_dict(record, config) for record in result.records]
    _write_csv(paths["detail"], detail_rows)
    _write_csv(paths["section_a"], [_record_to_dict(record, config) for record in result.section_a_books_not_in_2b])
    _write_csv(paths["section_b"], [_record_to_dict(record, config) for record in result.section_b_2b_not_in_books])
    _write_csv(paths["summary"], [_dataclass_to_dict(row) for row in result.monthly_summaries])

    open_items = [outstanding_to_dict(item) for item in result.outstanding_items]
    _atomic_json_write(paths["open_items"], {"outstanding_items": open_items})
    payload = {
        "as_of_month": result.as_of_month,
        "counts": _status_counts(result.records),
        "sections": {
            "books_not_in_gstr2b": [record.record_id for record in result.section_a_books_not_in_2b],
            "gstr2b_not_in_books": [record.record_id for record in result.section_b_2b_not_in_books],
        },
        "records": [_record_to_dict(record, config) for record in result.records],
        "monthly_summaries": [_dataclass_to_dict(row) for row in result.monthly_summaries],
        "outstanding_items": open_items,
    }
    _atomic_json_write(paths["json"], payload)
    write_excel_report(result, paths["workbook"], config)
    return paths


def _record_to_dict(record: ReconciliationRecord, config: ReconciliationConfig) -> dict[str, Any]:
    book = record.books
    gstr = record.gstr2b
    book_norm = normalized_values(book, config) if book else {}
    gstr_norm = normalized_values(gstr, config) if gstr else {}
    row: dict[str, Any] = {
        "record_id": record.record_id,
        "status": record.status,
        "match_level": record.match_level or "",
        "reason": record.reason,
        "resolved": record.resolved,
        "books_transaction_id": book.transaction_id if book else "",
        "gstr2b_transaction_id": gstr.transaction_id if gstr else "",
        "matched_transaction_id": gstr.transaction_id if book and gstr else "",
        "expected_gstr2b_month": record.expected_gstr2b_month or "",
        "books_month": book.books_month if book else "",
        "books_document_month": book.document_month if book else "",
        "document_month": (book.document_month if book else gstr.document_month if gstr else "") or "",
        "gstr2b_document_month": gstr.document_month if gstr else "",
        "gstr2b_month": gstr.gstr2b_month if gstr else "",
        "books_supplier_gstin": book.supplier_gstin if book else "",
        "gstr2b_supplier_gstin": gstr.supplier_gstin if gstr else "",
        "books_document_number": book.document_number if book else "",
        "books_document_number_aliases": ";".join(book.metadata.get("document_number_aliases", [])) if book else "",
        "gstr2b_document_number": gstr.document_number if gstr else "",
        "gstr2b_amends_document_number": gstr.amends_document_number if gstr else "",
        "gstr2b_is_amendment": gstr.is_amendment if gstr else "",
        "gstr2b_superseded_document_numbers": ";".join(gstr.metadata.get("superseded_document_numbers", [])) if gstr else "",
        "books_document_type": book.document_type if book else "",
        "gstr2b_document_type": gstr.document_type if gstr else "",
        "books_document_date": book.document_date if book else "",
        "gstr2b_document_date": gstr.document_date if gstr else "",
        "books_original_invoice_number": book.original_invoice_number if book else "",
        "gstr2b_original_invoice_number": gstr.original_invoice_number if gstr else "",
        "gstr2b_supplier_report_month": gstr.supplier_report_month if gstr else "",
        "books_supplier": book.supplier_name if book else "",
        "gstr2b_supplier": gstr.supplier_name if gstr else "",
        "books_classification": book.classification if book else "",
        "books_classification_reason": book.classification_reason if book else "",
        "books_source_file": book.source_file if book else "",
        "books_source_sheet": book.source_sheet if book else "",
        "books_source_row": book.source_row if book else "",
        "gstr2b_source_file": gstr.source_file if gstr else "",
        "gstr2b_source_sheet": gstr.source_sheet if gstr else "",
        "gstr2b_source_row": gstr.source_row if gstr else "",
        "gstr2b_month_basis": gstr.metadata.get("gstr2b_month_basis", "") if gstr else "",
        "gstr2b_itc_availability": gstr.itc_availability if gstr else "",
        "books_taxable_value": str(book.taxable_value) if book else "",
        "books_cgst": str(book.cgst) if book else "",
        "books_sgst": str(book.sgst) if book else "",
        "books_igst": str(book.igst) if book else "",
        "books_cess": str(book.cess) if book else "",
        "books_total_value": str(book.total_value) if book else "",
        "gstr2b_taxable_value": str(gstr.taxable_value) if gstr else "",
        "gstr2b_cgst": str(gstr.cgst) if gstr else "",
        "gstr2b_sgst": str(gstr.sgst) if gstr else "",
        "gstr2b_igst": str(gstr.igst) if gstr else "",
        "gstr2b_cess": str(gstr.cess) if gstr else "",
        "gstr2b_total_value": str(gstr.total_value) if gstr else "",
        "books_normalized_taxable_value": str(book_norm.get("taxable_value", "")),
        "books_normalized_cgst": str(book_norm.get("cgst", "")),
        "books_normalized_sgst": str(book_norm.get("sgst", "")),
        "books_normalized_igst": str(book_norm.get("igst", "")),
        "books_normalized_cess": str(book_norm.get("cess", "")),
        "gstr2b_normalized_taxable_value": str(gstr_norm.get("taxable_value", "")),
        "gstr2b_normalized_cgst": str(gstr_norm.get("cgst", "")),
        "gstr2b_normalized_sgst": str(gstr_norm.get("sgst", "")),
        "gstr2b_normalized_igst": str(gstr_norm.get("igst", "")),
        "gstr2b_normalized_cess": str(gstr_norm.get("cess", "")),
        "possible_match_ids": ";".join(record.possible_match_ids),
    }
    for name in ("taxable_value", "cgst", "sgst", "igst", "cess"):
        row[f"difference_{name}"] = str(record.differences.get(name, ""))
    return row


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    # Keep stable columns even for an empty section.
    keys: list[str] = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    if not keys:
        keys = ["record_id", "status", "reason"]
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_excel_report(
    result: ReconciliationResult,
    path: str | Path,
    config: ReconciliationConfig,
) -> Path:
    """Write a formatted multi-sheet XLSX deliverable."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError("Excel report generation requires openpyxl; install requirements.txt") from error

    workbook_path = Path(path)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Read Me"
    instructions.sheet_view.showGridLines = False
    instructions["A1"] = "GSTR-2B vs Books Reconciliation"
    instructions["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    instructions["A1"].fill = PatternFill("solid", fgColor="17365D")
    instructions.merge_cells("A1:B1")

    status_counts = _status_counts(result.records)
    mapping_bases = sorted({
        str(tx.metadata.get("gstr2b_month_basis", "unknown"))
        for record in result.records for tx in (record.gstr2b,) if tx is not None
    })
    facts = [
        ("As of month", result.as_of_month or "Not assigned"),
        ("Reconciliation records", len(result.records)),
        ("Open timing items", len(result.outstanding_items)),
        ("GSTR-2B month basis", ", ".join(mapping_bases) or "Not available"),
        (
            "Period caution",
            "The supplied GSTR-2B workbook aggregates periods. Filing-month is a proxy, not an authoritative 2B tax period. For monthly files, import with --gstr2b-period YYYY-MM.",
        ),
        (
            "Zero-tax classification",
            "Books exports lack an explicit ITC-eligibility field. Account labels are used first; otherwise zero GST with a GSTIN defaults to INELIGIBLE_ITC and without a GSTIN to NON_GST. Confirm/configure this policy.",
        ),
        (
            "Matching/sign convention",
            "Amounts are normalized before comparison. Credit notes reduce ITC; the supplied Books Debit Note register maps to a supplier Credit Note by default. See config.example.json to change aliases/signs/tolerances.",
        ),
        (
            "Month fields",
            "Books month, document month and GSTR-2B month are separate. A cross-month match is one timing-difference record and is carried/cleared through the outstanding register.",
        ),
    ]
    for row_index, (label, value) in enumerate(facts, start=3):
        instructions.cell(row_index, 1, label).font = Font(bold=True, color="17365D")
        instructions.cell(row_index, 2, value).alignment = Alignment(vertical="top", wrap_text=True)
    row_index = len(facts) + 4
    instructions.cell(row_index, 1, "Status counts").font = Font(bold=True, color="17365D")
    for offset, (status, count) in enumerate(sorted(status_counts.items()), start=1):
        instructions.cell(row_index + offset, 1, status)
        instructions.cell(row_index + offset, 2, count)
    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 112
    instructions.freeze_panes = "A3"

    detail_rows = [_record_to_dict(record, config) for record in result.records]
    detail_headers = list(detail_rows[0]) if detail_rows else ["record_id", "status", "reason"]
    monthly_rows = [_dataclass_to_dict(row) for row in result.monthly_summaries]
    monthly_headers = list(monthly_rows[0]) if monthly_rows else ["month"]
    open_rows = [_outstanding_report_row(item) for item in result.outstanding_items]
    open_headers = list(open_rows[0]) if open_rows else [
        "id", "supplier_gstin", "document_number", "document_type", "books_month",
        "gstr2b_month", "expected_gstr2b_month", "status", "matched_transaction_id",
    ]

    _add_excel_sheet(workbook, "Monthly Summary", monthly_headers, monthly_rows, "MonthlySummary", get_column_letter, Table, TableStyleInfo, Font, PatternFill, Alignment)
    _add_excel_sheet(workbook, "Reconciliation", detail_headers, detail_rows, "ReconDetail", get_column_letter, Table, TableStyleInfo, Font, PatternFill, Alignment)
    _add_excel_sheet(
        workbook, "Books Not in 2B", detail_headers,
        [_record_to_dict(record, config) for record in result.section_a_books_not_in_2b],
        "BooksNotIn2B", get_column_letter, Table, TableStyleInfo, Font, PatternFill, Alignment,
    )
    _add_excel_sheet(
        workbook, "2B Not in Books", detail_headers,
        [_record_to_dict(record, config) for record in result.section_b_2b_not_in_books],
        "TwoBNotInBooks", get_column_letter, Table, TableStyleInfo, Font, PatternFill, Alignment,
    )
    _add_excel_sheet(workbook, "Outstanding", open_headers, open_rows, "OutstandingItems", get_column_letter, Table, TableStyleInfo, Font, PatternFill, Alignment)
    workbook.properties.title = "GSTR-2B vs Books Reconciliation"
    workbook.properties.subject = "Month-wise purchase and ITC reconciliation"
    workbook.properties.creator = "GSTR-2B Reconciliation"
    workbook.active = 0
    workbook.save(workbook_path)
    return workbook_path


def _outstanding_report_row(item: OutstandingItem) -> dict[str, Any]:
    tx = item.transaction
    return {
        "id": item.id,
        "supplier_gstin": item.supplier_gstin,
        "document_number": item.document_number,
        "document_type": item.document_type,
        "document_date": item.document_date,
        "books_month": item.books_month,
        "document_month": tx.document_month,
        "gstr2b_month": item.gstr2b_month,
        "expected_gstr2b_month": item.expected_gstr2b_month,
        "taxable_value": str(item.taxable_value),
        "cgst": str(item.cgst),
        "sgst": str(item.sgst),
        "igst": str(item.igst),
        "cess": str(item.cess),
        "normalized_taxable_value": str(item.normalized_taxable_value),
        "normalized_cgst": str(item.normalized_cgst),
        "normalized_sgst": str(item.normalized_sgst),
        "normalized_igst": str(item.normalized_igst),
        "normalized_cess": str(item.normalized_cess),
        "status": item.status,
        "matched_transaction_id": item.matched_transaction_id or "",
        "source_file": tx.source_file,
        "source_sheet": tx.source_sheet,
        "source_row": tx.source_row or "",
    }


def _add_excel_sheet(
    workbook: Any,
    title: str,
    headers: list[str],
    rows: list[dict[str, Any]],
    table_name: str,
    get_column_letter: Any,
    Table: Any,
    TableStyleInfo: Any,
    Font: Any,
    PatternFill: Any,
    Alignment: Any,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(header, row.get(header)) for header in headers])

    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    sheet.auto_filter.ref = sheet.dimensions

    if rows:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    for column_index, header in enumerate(headers, start=1):
        label = header.lower()
        is_currency = any(token in label for token in (
            "taxable_value", "total_value", "_cgst", "_sgst", "_igst", "_cess",
            "_amount", "_itc", "difference_", "timing_difference",
        )) or label in {"cgst", "sgst", "igst", "cess", "taxable_value"}
        if is_currency:
            for row_index in range(2, sheet.max_row + 1):
                sheet.cell(row_index, column_index).number_format = '#,##0.00;[Red](#,##0.00)'
        sample_lengths = [
            len(str(sheet.cell(row_index, column_index).value or ""))
            for row_index in range(1, min(sheet.max_row, 100) + 1)
        ]
        width = min(max(max(sample_lengths, default=0) + 2, 12), 44)
        if "reason" in label or "classification_reason" in label:
            width = 52
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _excel_value(header: str, value: Any) -> Any:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    label = header.lower()
    is_currency = any(token in label for token in (
        "taxable_value", "total_value", "_cgst", "_sgst", "_igst", "_cess",
        "_amount", "_itc", "difference_", "timing_difference",
    )) or label in {"cgst", "sgst", "igst", "cess", "taxable_value"}
    if is_currency or label in {"source_row"}:
        try:
            number = Decimal(str(value))
            if number.is_finite():
                return int(number) if label == "source_row" and number == number.to_integral_value() else float(number)
        except (InvalidOperation, TypeError, ValueError):
            pass
    return value


def _dataclass_to_dict(value: Any) -> dict[str, Any]:
    return {
        item.name: str(getattr(value, item.name)) if isinstance(getattr(value, item.name), Decimal) else getattr(value, item.name)
        for item in fields(value)
    }


def _status_counts(records: list[ReconciliationRecord]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for record in records:
        counts[record.status] = counts.get(record.status, 0) + 1
    return counts


def _atomic_json_write(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2, default=_json_default)
            stream.write("\n")
        os.replace(temporary_name, path)
    except Exception:
        try:
            os.unlink(temporary_name)
        except OSError:
            pass
        raise


def _json_default(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")


def _decimal(value: Any) -> Decimal:
    return value if isinstance(value, Decimal) else Decimal(str(value or 0))
