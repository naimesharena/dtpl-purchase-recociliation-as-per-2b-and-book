#!/usr/bin/env python3
"""Build a monthly GSTR-2B vs books reconciliation workbook.

Inputs are the seven source workbooks in this folder. Six of the books registers
have an .xls suffix but are OOXML workbooks; the GSTR-2B download is a legacy
BIFF .xls file. Required packages: xlrd, openpyxl, XlsxWriter.

Run: python build_reconciliation.py
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import sys

import openpyxl
import xlrd
import xlsxwriter

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "GSTR2B_Books_Reconciliation_FY2025-26.xlsx"
PURCHASE_FILES = ["PURCHASE-HR.xls", "PURCHASE-PL.xls", "PURCHASE-VL.xls"]
DEBIT_NOTE_FILES = ["DEBIT NOTE-HR.xls", "DEBIT NOTE-PL.xls", "DEBIT NOTE-VL.xls"]
GSTR_FILE = "gstr-2B.xls"
FY_MONTHS = [datetime(2025 + ((4 + i - 1) // 12), ((4 + i - 1) % 12) + 1, 1) for i in range(12)]
TOLERANCE = 0.02


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value) -> float:
    if isinstance(value, bool) or value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(",", "").replace("₹", "")
    if raw in ("-", "—", "–"):
        return 0.0
    try:
        return float(raw)
    except ValueError:
        return 0.0


def norm_doc(value) -> str:
    """Normalize reference numbers conservatively: uppercase, strip punctuation."""
    return re.sub(r"[^A-Z0-9]", "", text(value).upper())


def norm_gstin(value) -> str:
    return re.sub(r"\s+", "", text(value).upper())


def parse_source_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    raw = text(value)
    if not raw:
        return None
    for fmt in ("%d-%m-%Y", "%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def period_date(value):
    raw = text(value)
    match = re.fullmatch(r"(\d{2})(\d{4})", raw)
    if not match:
        return None
    month, year = int(match.group(1)), int(match.group(2))
    if not 1 <= month <= 12:
        return None
    return datetime(year, month, 1)


def month_key(value):
    if not value:
        return ""
    return value.strftime("%Y-%m")


def amount_sum(row, indices):
    return sum(number(row[i]) for i in indices)


def amount_paise(value):
    """Return a currency amount as integer paise for strict amount-key matching."""
    return int(round(number(value) * 100))


def has_tax(record) -> bool:
    return abs(record["cgst"]) + abs(record["sgst"]) + abs(record["igst"]) + abs(record.get("cess", 0.0)) > 0.005


def bucket_for(availability, rcm, rejected=False):
    if rejected:
        return "Rejected"
    availability = text(availability).lower()
    is_rcm = text(rcm).upper() == "Y"
    if availability == "yes":
        return "RCM available" if is_rcm else "Available"
    if availability == "no":
        return "RCM not available" if is_rcm else "Not available"
    return "Unknown"


def _book_col_map(headers):
    norm_headers = [text(h).lower().replace("_x0004_", "") for h in headers]
    def find_contains(needles, default=None):
        for i, h in enumerate(norm_headers):
            if any(n in h for n in needles):
                return i
        return default
    date_i = find_contains(["date"], 0)
    party_i = find_contains(["particulars"], 1)
    voucher_type_i = find_contains(["voucher type"], 2)
    voucher_i = find_contains(["voucher no"], 3)
    if voucher_i is None:
        voucher_i = 3
    inv_i = find_contains(["supplier invoice no"], 4)
    inv_date_i = find_contains(["supplier invoice date", "voucher ref. date"], 5)
    gstin_i = find_contains(["gstin/uin", "gstin"], 6)
    gross_i = find_contains(["gross total"], 9)
    round_i = find_contains(["round off"], None)
    all_items_i = find_contains(["all items"], None)
    purchase_accounts_i = find_contains(["purchase accounts"], None)
    tax_cols = {"cgst": [], "sgst": [], "igst": [], "cess": []}
    for i, h in enumerate(norm_headers):
        if "cgst" in h and ("tax" in h or "gst" in h):
            tax_cols["cgst"].append(i)
        elif "sgst" in h and ("tax" in h or "gst" in h):
            tax_cols["sgst"].append(i)
        elif "igst" in h and ("tax" in h or "gst" in h):
            tax_cols["igst"].append(i)
        elif "cess" in h and ("tax" in h or "cess" in h):
            tax_cols["cess"].append(i)
    # In these Tally registers, all posting/account columns start immediately
    # after Gross Total. Tax ledgers, round-off, and the special All Items field
    # are excluded from the book base.
    excluded = set(sum(tax_cols.values(), []))
    if round_i is not None:
        excluded.add(round_i)
    if all_items_i is not None:
        excluded.add(all_items_i)
    account_cols = [i for i in range((gross_i or 9) + 1, len(headers)) if i not in excluded]
    return {
        "date": date_i, "party": party_i, "voucher_type": voucher_type_i,
        "voucher": voucher_i, "invoice": inv_i, "invoice_date": inv_date_i,
        "gstin": gstin_i, "gross": gross_i, "round": round_i,
        "all_items": all_items_i, "purchase_accounts": purchase_accounts_i,
        "tax_cols": tax_cols, "account_cols": account_cols,
    }


def read_book_register(filename, kind):
    path = ROOT / filename
    data = path.read_bytes()
    if data[:2] != b"PK":
        raise ValueError(f"Expected an OOXML workbook with .xls suffix: {filename}")
    wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=10, max_row=10))]
    cols = _book_col_map(headers)
    results = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=11, values_only=True), start=11):
        if not row or all(v is None for v in row):
            continue
        posting_date = parse_source_date(row[cols["date"]] if cols["date"] is not None else None)
        party = text(row[cols["party"]] if cols["party"] is not None else "")
        if not posting_date or party.lower() == "grand total":
            continue
        raw_accounts = amount_sum(row, cols["account_cols"])
        cgst = amount_sum(row, cols["tax_cols"]["cgst"])
        sgst = amount_sum(row, cols["tax_cols"]["sgst"])
        igst = amount_sum(row, cols["tax_cols"]["igst"])
        cess = amount_sum(row, cols["tax_cols"]["cess"])
        gross = number(row[cols["gross"]] if cols["gross"] is not None else 0)
        round_off = number(row[cols["round"]] if cols["round"] is not None else 0)
        all_items = number(row[cols["all_items"]] if cols["all_items"] is not None else 0)
        purchase_accounts = number(row[cols["purchase_accounts"]] if cols["purchase_accounts"] is not None else 0)
        total_tax = cgst + sgst + igst + cess
        # The source register's Gross Total can be net of All Items/withholding.
        # Use the invoice-total bridge for the reconciliation base. A small set
        # of Purchase Register rows has a contra/withholding amount in the
        # Purchase Accounts group: when the account sum exceeds this bridge by
        # exactly twice that source amount, add it back once (the account sum
        # then equals base + the withheld amount). Keep the ledger sum and QA
        # difference on every output row rather than silently losing this issue.
        base_bridge = gross + all_items - total_tax - round_off
        special_purchase_account_adjustment = (
            purchase_accounts
            if kind == "Purchase"
            and abs(purchase_accounts) > 0.005
            and abs((raw_accounts - base_bridge) - 2 * purchase_accounts) <= 0.05
            else 0.0
        )
        book_base = base_bridge + special_purchase_account_adjustment
        base_method = "Gross + All Items − GST − Round Off"
        if abs(special_purchase_account_adjustment) > 0.005:
            base_method = "Gross + Purchase Accounts adjustment + All Items − GST − Round Off"
        source_file = filename
        branch_match = re.search(r"-(HR|PL|VL)\.xls$", filename, re.I)
        branch = branch_match.group(1).upper() if branch_match else ""
        voucher = text(row[cols["voucher"]] if cols["voucher"] is not None else "")
        invoice = text(row[cols["invoice"]] if cols["invoice"] is not None else "")
        invoice_date = parse_source_date(row[cols["invoice_date"]] if cols["invoice_date"] is not None else None)
        # For the debit note register the reference date is the note date; if it
        # is missing, use the book posting date for duplicate-key resolution.
        doc_date = invoice_date or posting_date
        if kind == "Debit Note":
            doc_date = invoice_date or posting_date
        record = {
            "id": f"{Path(filename).stem}::{excel_row}",
            "kind": kind,
            "source_file": source_file,
            "source_row": excel_row,
            "branch": branch,
            "date": posting_date,
            "month": month_key(posting_date),
            "party": party,
            "voucher_type": text(row[cols["voucher_type"]] if cols["voucher_type"] is not None else ""),
            "voucher": voucher,
            "invoice": invoice,
            "doc_date": doc_date,
            "gstin": norm_gstin(row[cols["gstin"]] if cols["gstin"] is not None else ""),
            "gross": gross,
            "base": book_base,
            "ledger_sum": raw_accounts,
            "base_qa_delta": raw_accounts - book_base,
            "base_method": base_method,
            "base_bridge": base_bridge,
            "purchase_accounts_source": purchase_accounts,
            "purchase_account_adjustment": special_purchase_account_adjustment,
            "cgst": cgst,
            "sgst": sgst,
            "igst": igst,
            "cess": cess,
            "gst": total_tax,
            "round_off": round_off,
            "all_items": all_items,
            "no_tax": not (abs(cgst) + abs(sgst) + abs(igst) + abs(cess) > 0.005),
            "headers": headers,
        }
        record["signed"] = -1 if kind == "Debit Note" else 1
        record["signed_base"] = record["signed"] * record["base"]
        record["signed_cgst"] = record["signed"] * cgst
        record["signed_sgst"] = record["signed"] * sgst
        record["signed_igst"] = record["signed"] * igst
        record["signed_gst"] = record["signed"] * record["gst"]
        record["no_tax_base_signed"] = record["signed_base"] if record["no_tax"] else 0.0
        record["taxed_base_signed"] = 0.0 if record["no_tax"] else record["signed_base"]
        results.append(record)
    wb.close()
    return results


def read_gstr_workbook():
    path = ROOT / GSTR_FILE
    wb = xlrd.open_workbook(str(path), on_demand=True)
    return wb


def extract_gstr_records(wb):
    b2b, notes, amendments, rejected = [], [], [], []
    sheet_specs = []

    # B2B invoices (original detail rows start at Excel row 7).
    ws = wb.sheet_by_name("B2B")
    for i in range(6, ws.nrows):
        r = ws.row_values(i)
        if not text(r[0]) or not text(r[2]):
            continue
        rec = {
            "id": f"B2B::{i+1}", "source_sheet": "B2B", "source_row": i + 1,
            "event": "Original", "transaction": "Invoice", "gstin": norm_gstin(r[0]),
            "supplier": text(r[1]), "doc_no": text(r[2]), "doc_type": text(r[3]),
            "doc_date": parse_source_date(r[4]), "invoice_value": number(r[5]),
            "pos": text(r[6]), "rcm": text(r[7]).upper(), "base": number(r[8]),
            "igst": number(r[9]), "cgst": number(r[10]), "sgst": number(r[11]),
            "cess": number(r[12]), "period_raw": text(r[13]), "period": period_date(r[13]),
            "filing_date": parse_source_date(r[14]), "availability": text(r[15]),
            "reason": text(r[16]), "rate": text(r[17]), "note_type": "",
            "reference_doc_no": text(r[2]), "raw_sign": 1,
        }
        rec["bucket"] = bucket_for(rec["availability"], rec["rcm"])
        rec["signed_base"] = rec["base"]
        for comp in ("igst", "cgst", "sgst", "cess"):
            rec[f"signed_{comp}"] = rec[comp]
        rec["signed_gst"] = rec["signed_igst"] + rec["signed_cgst"] + rec["signed_sgst"] + rec["signed_cess"]
        b2b.append(rec)

    # Original supplier debit/credit notes. C = recipient input-credit reduction;
    # D = recipient input-credit increase.
    ws = wb.sheet_by_name("B2B-CDNR")
    for i in range(6, ws.nrows):
        r = ws.row_values(i)
        if not text(r[0]) or not text(r[2]):
            continue
        note_type = text(r[3]).upper()
        sign = -1 if note_type == "C" else 1
        rec = {
            "id": f"B2B-CDNR::{i+1}", "source_sheet": "B2B-CDNR", "source_row": i + 1,
            "event": "Original", "transaction": "Credit note" if note_type == "C" else "Debit note",
            "gstin": norm_gstin(r[0]), "supplier": text(r[1]), "doc_no": text(r[2]),
            "doc_type": note_type, "doc_date": parse_source_date(r[5]),
            "invoice_value": number(r[6]), "pos": text(r[7]), "rcm": text(r[8]).upper(),
            "base": number(r[9]), "igst": number(r[10]), "cgst": number(r[11]),
            "sgst": number(r[12]), "cess": number(r[13]), "period_raw": text(r[14]),
            "period": period_date(r[14]), "filing_date": parse_source_date(r[15]),
            "availability": text(r[16]), "reason": text(r[17]), "rate": text(r[18]),
            "note_type": note_type, "reference_doc_no": text(r[2]), "raw_sign": sign,
        }
        rec["bucket"] = bucket_for(rec["availability"], rec["rcm"])
        rec["signed_base"] = sign * rec["base"]
        for comp in ("igst", "cgst", "sgst", "cess"):
            rec[f"signed_{comp}"] = sign * rec[comp]
        rec["signed_gst"] = rec["signed_igst"] + rec["signed_cgst"] + rec["signed_sgst"] + rec["signed_cess"]
        notes.append(rec)

    # Invoice amendments: retain the original month, then record only the change
    # in the supplier's amended filing period to prevent counting the revised
    # amount a second time.
    ws = wb.sheet_by_name("B2BA")
    for i in range(7, ws.nrows):
        r = ws.row_values(i)
        if not text(r[2]) or not text(r[4]):
            continue
        old_no, old_date, gstin, supplier = text(r[0]), parse_source_date(r[1]), norm_gstin(r[2]), text(r[3])
        new = {
            "doc_no": text(r[4]), "doc_type": text(r[5]), "doc_date": parse_source_date(r[6]),
            "invoice_value": number(r[7]), "pos": text(r[8]), "rcm": text(r[9]).upper(),
            "base": number(r[10]), "igst": number(r[11]), "cgst": number(r[12]),
            "sgst": number(r[13]), "cess": number(r[14]), "period_raw": text(r[15]),
            "period": period_date(r[15]), "filing_date": parse_source_date(r[16]),
            "availability": text(r[17]), "reason": text(r[18]), "rate": text(r[19]),
        }
        rec = {
            "id": f"B2BA::{i+1}", "source_sheet": "B2BA", "source_row": i + 1,
            "event": "Amendment adjustment", "transaction": "Invoice amendment",
            "gstin": gstin, "supplier": supplier, "old_doc_no": old_no,
            "old_doc_date": old_date, "doc_no": new["doc_no"], "doc_type": new["doc_type"],
            "doc_date": new["doc_date"], "invoice_value": new["invoice_value"],
            "pos": new["pos"], "rcm": new["rcm"], "base": new["base"],
            "igst": new["igst"], "cgst": new["cgst"], "sgst": new["sgst"],
            "cess": new["cess"], "period_raw": new["period_raw"], "period": new["period"],
            "filing_date": new["filing_date"], "availability": new["availability"],
            "reason": new["reason"], "rate": new["rate"], "note_type": "",
            "reference_doc_no": new["doc_no"], "raw_sign": 1,
        }
        rec["bucket"] = bucket_for(rec["availability"], rec["rcm"])
        amendments.append(rec)

    # Credit/debit-note amendments.
    ws = wb.sheet_by_name("B2B-CDNRA")
    for i in range(7, ws.nrows):
        r = ws.row_values(i)
        if not text(r[3]) or not text(r[5]):
            continue
        old_type, old_no = text(r[0]).upper(), text(r[1])
        gstin, supplier = norm_gstin(r[3]), text(r[4])
        note_type = text(r[6]).upper()
        sign = -1 if note_type == "C" else 1
        rec = {
            "id": f"B2B-CDNRA::{i+1}", "source_sheet": "B2B-CDNRA", "source_row": i + 1,
            "event": "Amendment adjustment", "transaction": "Credit note amendment" if note_type == "C" else "Debit note amendment",
            "gstin": gstin, "supplier": supplier, "old_doc_no": old_no,
            "old_doc_type": old_type, "old_doc_date": parse_source_date(r[2]),
            "doc_no": text(r[5]), "doc_type": note_type, "doc_date": parse_source_date(r[8]),
            "invoice_value": number(r[9]), "pos": text(r[10]), "rcm": text(r[11]).upper(),
            "base": number(r[12]), "igst": number(r[13]), "cgst": number(r[14]),
            "sgst": number(r[15]), "cess": number(r[16]), "period_raw": text(r[17]),
            "period": period_date(r[17]), "filing_date": parse_source_date(r[18]),
            "availability": text(r[19]), "reason": text(r[20]), "rate": text(r[21]),
            "note_type": note_type, "reference_doc_no": text(r[5]), "raw_sign": sign,
        }
        rec["bucket"] = bucket_for(rec["availability"], rec["rcm"])
        amendments.append(rec)

    # IMS-rejected invoices are on a separate sheet and are not eligible ITC.
    # This sheet's data is offset by one column relative to its visible header:
    # taxable base is in col I (index 8), tax components are J:M (9:12).
    ws = wb.sheet_by_name("B2B(Rejected)")
    for i in range(6, ws.nrows):
        r = ws.row_values(i)
        if not text(r[0]) or not text(r[2]):
            continue
        rec = {
            "id": f"B2B(Rejected)::{i+1}", "source_sheet": "B2B(Rejected)", "source_row": i + 1,
            "event": "Rejected", "transaction": "Rejected invoice", "gstin": norm_gstin(r[0]),
            "supplier": text(r[1]), "doc_no": text(r[2]), "doc_type": text(r[3]),
            "doc_date": parse_source_date(r[4]), "invoice_value": number(r[5]),
            "pos": text(r[6]), "rcm": "N", "base": number(r[8]), "igst": number(r[9]),
            "cgst": number(r[10]), "sgst": number(r[11]), "cess": number(r[12]),
            "period_raw": text(r[13]), "period": period_date(r[13]),
            "filing_date": parse_source_date(r[14]), "availability": "Rejected",
            "reason": "IMS rejected record", "rate": text(r[17]), "note_type": "",
            "reference_doc_no": text(r[2]), "raw_sign": 1,
        }
        rec["bucket"] = "Rejected"
        rec["signed_base"] = rec["base"]
        for comp in ("igst", "cgst", "sgst", "cess"):
            rec[f"signed_{comp}"] = rec[comp]
        rec["signed_gst"] = rec["signed_igst"] + rec["signed_cgst"] + rec["signed_sgst"] + rec["signed_cess"]
        rejected.append(rec)

    return b2b, notes, amendments, rejected


def pair_score(book, gstr, note=False):
    book_tax = book["cgst"] + book["sgst"] + book["igst"] + book.get("cess", 0.0)
    gstr_tax = gstr["cgst"] + gstr["sgst"] + gstr["igst"] + gstr["cess"]
    base_diff = abs(book["base"] - gstr["base"])
    tax_diff = abs(book_tax - gstr_tax)
    bdate = book.get("doc_date") or book.get("date")
    gdate = gstr.get("doc_date")
    date_diff = abs((bdate - gdate).days) if bdate and gdate else 999999
    # Value similarity comes first, then document-date proximity. This scoring
    # is used only to resolve duplicate references; unique GSTIN+reference
    # matches are retained even when their amount differs.
    return (round(base_diff + tax_diff, 2), date_diff, round(base_diff, 2), round(tax_diff, 2), book["id"])


def match_one_to_one(book_records, gstr_records, primary_field, alternative_fields,
                     uniqueness_book_records=None, uniqueness_gstr_records=None):
    """Match exact GSTIN/reference first, then only unique cross-GSTIN ref/base pairs."""
    if uniqueness_book_records is None:
        uniqueness_book_records = book_records
    if uniqueness_gstr_records is None:
        uniqueness_gstr_records = gstr_records
    matches_by_gstr = {}
    matches_by_book = {}
    duplicate_groups = {}
    all_candidate_book_ids = defaultdict(set)
    used_books = set()

    def records_for_field(field):
        index = defaultdict(list)
        for b in book_records:
            token = norm_doc(b.get(field))
            if b.get("gstin") and token:
                index[(b["gstin"], token)].append(b)
        return index

    primary_index = records_for_field(primary_field)
    alt_indices = [(field, records_for_field(field)) for field in alternative_fields]

    for g in gstr_records:
        token = norm_doc(g.get("doc_no"))
        if not g.get("gstin") or not token:
            continue
        key = (g["gstin"], token)
        for b in primary_index.get(key, []):
            all_candidate_book_ids[g["id"]].add(b["id"])
        for _, idx in alt_indices:
            for b in idx.get(key, []):
                all_candidate_book_ids[g["id"]].add(b["id"])

    def apply_field(field, index, method_name):
        groups = defaultdict(list)
        for g in gstr_records:
            if g["id"] in matches_by_gstr:
                continue
            token = norm_doc(g.get("doc_no"))
            if not g.get("gstin") or not token:
                continue
            key = (g["gstin"], token)
            candidates = [b for b in index.get(key, []) if b["id"] not in used_books]
            if candidates:
                groups[key].append((g, candidates))
        for key, g_groups in groups.items():
            # The source exports in this folder have unique GSTR references. If
            # future data has repeated GSTR keys, process one-to-one by the same
            # amount/date score and flag the group for review.
            available_g = [g for g, _ in g_groups]
            book_candidates = list({b["id"]: b for _, cands in g_groups for b in cands}.values())
            if len(available_g) == 1:
                g = available_g[0]
                cands = [b for b in book_candidates if b["id"] not in used_books]
                if not cands:
                    continue
                cands.sort(key=lambda b: pair_score(b, g, note=(g.get("transaction", "").lower().find("note") >= 0)))
                selected = cands[0]
                duplicate = len(cands) > 1
                tied = duplicate and pair_score(cands[0], g)[:-1] == pair_score(cands[1], g)[:-1]
                matches_by_gstr[g["id"]] = {
                    "book": selected, "method": method_name,
                    "duplicate": duplicate, "ambiguous": tied, "gstin_mismatch": False,
                }
                matches_by_book[selected["id"]] = g
                used_books.add(selected["id"])
                if duplicate:
                    duplicate_groups.setdefault(key, set()).update(b["id"] for b in cands)
            else:
                remaining_g = list(available_g)
                remaining_b = [b for b in book_candidates if b["id"] not in used_books]
                while remaining_g and remaining_b:
                    options = []
                    for g in remaining_g:
                        for b in remaining_b:
                            options.append((pair_score(b, g), g, b))
                    options.sort(key=lambda x: x[0])
                    best_score, g, b = options[0]
                    matches_by_gstr[g["id"]] = {"book": b, "method": method_name, "duplicate": True, "ambiguous": len(options) > 1 and options[0][0][:-1] == options[1][0][:-1], "gstin_mismatch": False}
                    matches_by_book[b["id"]] = g
                    used_books.add(b["id"])
                    duplicate_groups.setdefault(key, set()).update(x["id"] for x in remaining_b)
                    remaining_g.remove(g)
                    remaining_b.remove(b)

    apply_field(primary_field, primary_index, "Voucher No. ↔ 2B document number")
    for field, idx in alt_indices:
        method = "Supplier invoice number ↔ 2B document number" if field == "invoice" else f"{field} ↔ 2B document number"
        apply_field(field, idx, method)

    # Conservative fallback for a supplier-registration mismatch: only link an
    # otherwise-unmatched pair when the normalized document reference AND the
    # taxable base (rounded to paise) are unique on both complete source sides.
    # Exact GSTIN matches above always take precedence. Ambiguous candidates are
    # left open and flagged for review rather than force-matched.
    book_by_id = {b["id"]: b for b in book_records}
    gstr_by_id = {g["id"]: g for g in gstr_records}
    all_book_by_id = {b["id"]: b for b in uniqueness_book_records}
    all_book_by_id.update(book_by_id)
    all_gstr_by_id = {g["id"]: g for g in uniqueness_gstr_records}
    all_gstr_by_id.update(gstr_by_id)
    book_ids_by_ref_amount = defaultdict(set)
    gstr_ids_by_ref_amount = defaultdict(set)
    for b in uniqueness_book_records:
        amount_key = amount_paise(b.get("base", 0.0))
        for field in (primary_field, *alternative_fields):
            token = norm_doc(b.get(field))
            if token:
                book_ids_by_ref_amount[(token, amount_key)].add(b["id"])
    for g in uniqueness_gstr_records:
        token = norm_doc(g.get("doc_no"))
        if token:
            gstr_ids_by_ref_amount[(token, amount_paise(g.get("base", 0.0)))].add(g["id"])

    cross_candidates_by_book = defaultdict(set)
    cross_candidates_by_gstr = defaultdict(set)
    fallback_method = "Unique document reference + taxable base (paise match); GSTIN differs"
    for key in sorted(set(book_ids_by_ref_amount) & set(gstr_ids_by_ref_amount)):
        book_ids = book_ids_by_ref_amount[key]
        gstr_ids = gstr_ids_by_ref_amount[key]
        cross_pairs = [
            (bid, gid) for bid in book_ids for gid in gstr_ids
            if all_book_by_id[bid].get("gstin")
            and all_gstr_by_id[gid].get("gstin")
            and all_book_by_id[bid]["gstin"] != all_gstr_by_id[gid]["gstin"]
        ]
        if not cross_pairs:
            continue
        if len(book_ids) == 1 and len(gstr_ids) == 1:
            bid, gid = next(iter(book_ids)), next(iter(gstr_ids))
            b, g = all_book_by_id[bid], all_gstr_by_id[gid]
            if bid in book_by_id and gid in gstr_by_id and bid not in used_books and gid not in matches_by_gstr:
                matches_by_gstr[gid] = {
                    "book": b, "method": fallback_method,
                    "duplicate": False, "ambiguous": False, "gstin_mismatch": True,
                }
                matches_by_book[bid] = g
                used_books.add(bid)
            else:
                if bid in book_by_id and bid not in used_books:
                    cross_candidates_by_book[bid].add(gid)
                if gid in gstr_by_id and gid not in matches_by_gstr:
                    cross_candidates_by_gstr[gid].add(bid)
        else:
            # At least one complete source side has a duplicate ref/base key.
            # Do not choose among candidates, even if a score would pick one.
            for bid, gid in cross_pairs:
                if bid in book_by_id and bid not in used_books:
                    cross_candidates_by_book[bid].add(gid)
                if gid in gstr_by_id and gid not in matches_by_gstr:
                    cross_candidates_by_gstr[gid].add(bid)

    # Mark leftover book references which point at a 2B reference already used
    # by a different book row; these are likely duplicate/reused references.
    candidate_to_books = defaultdict(set)
    for g_id, candidate_ids in all_candidate_book_ids.items():
        for bid in candidate_ids:
            candidate_to_books[g_id].add(bid)
    for b in book_records:
        b["match_info"] = None
        b["possible_duplicate"] = False
        b["candidate_gstr_ids"] = []
        b["cross_gstin_candidate_ids"] = []
        b["cross_gstin_review"] = False
    book_by_id = {b["id"]: b for b in book_records}
    for g in gstr_records:
        g["cross_gstin_candidate_book_ids"] = []
        g["cross_gstin_review"] = False
        info = matches_by_gstr.get(g["id"])
        if info:
            b = info["book"]
            b["match_info"] = {"gstr": g, "method": info["method"], "duplicate": info["duplicate"], "ambiguous": info["ambiguous"], "gstin_mismatch": info.get("gstin_mismatch", False)}
        for bid in all_candidate_book_ids.get(g["id"], set()):
            if bid in book_by_id:
                book_by_id[bid]["candidate_gstr_ids"].append(g["id"])
    for bid, gids in cross_candidates_by_book.items():
        b = book_by_id.get(bid)
        if b and not b["match_info"]:
            b["cross_gstin_candidate_ids"] = sorted(gids)
            b["cross_gstin_review"] = True
    for gid, bids in cross_candidates_by_gstr.items():
        g = gstr_by_id.get(gid)
        if g and gid not in matches_by_gstr:
            g["cross_gstin_candidate_book_ids"] = sorted(bids)
            g["cross_gstin_review"] = True
    for b in book_records:
        if b["candidate_gstr_ids"] and not b["match_info"]:
            b["possible_duplicate"] = True
        elif b["match_info"] and b["match_info"]["duplicate"]:
            b["possible_duplicate"] = True
        if b["cross_gstin_review"]:
            b["possible_duplicate"] = True
    return matches_by_gstr, matches_by_book, used_books


def signed_components(rec, sign=None):
    s = rec.get("raw_sign", 1) if sign is None else sign
    return {
        "base": s * rec["base"], "igst": s * rec["igst"],
        "cgst": s * rec["cgst"], "sgst": s * rec["sgst"],
        "cess": s * rec["cess"],
    }


def bucket_components(bucket, comp):
    names = ("base", "igst", "cgst", "sgst", "cess")
    result = {b: {k: 0.0 for k in names} for b in ("Available", "RCM available", "Not available", "RCM not available", "Rejected", "Unknown")}
    if bucket in result:
        result[bucket] = dict(comp)
    return result


def find_original_for_amendment(amend, originals):
    old_no = amend.get("old_doc_no", "")
    gstin = amend.get("gstin", "")
    target = norm_doc(old_no or amend.get("doc_no"))
    candidates = [r for r in originals if r["gstin"] == gstin and norm_doc(r["doc_no"]) == target]
    if len(candidates) == 1:
        return candidates[0], False
    if len(candidates) > 1:
        candidates.sort(key=lambda x: pair_score({"base": x["base"], "cgst": x["cgst"], "sgst": x["sgst"], "igst": x["igst"], "cess": x["cess"], "doc_date": x.get("doc_date"), "date": x.get("doc_date"), "id": x["id"]}, amend))
        return candidates[0], True
    return None, False


def normalise_amendments(amendments, invoice_originals, note_originals):
    normalized = []
    for amend in amendments:
        if amend["source_sheet"] == "B2BA":
            old, ambiguous = find_original_for_amendment(amend, invoice_originals)
        else:
            old, ambiguous = find_original_for_amendment(amend, note_originals)
        amend["original"] = old
        amend["old_lookup_ambiguous"] = ambiguous
        amend["old_found"] = old is not None
        new_signed = signed_components(amend)
        old_signed = signed_components(old) if old else {k: 0.0 for k in new_signed}
        # The old and revised availability/RCM categories may differ. The
        # bucket deltas below remove the old version from its bucket and add the
        # revised version to its new bucket.
        old_bucket = old["bucket"] if old else "Unknown"
        new_bucket = amend["bucket"]
        bucket_delta = {b: {k: 0.0 for k in new_signed} for b in ("Available", "RCM available", "Not available", "RCM not available", "Rejected", "Unknown")}
        if old:
            for k in new_signed:
                bucket_delta[old_bucket][k] -= old_signed[k]
        for k in new_signed:
            bucket_delta[new_bucket][k] += new_signed[k]
        delta = {k: new_signed[k] - old_signed[k] for k in new_signed}
        amend["old_bucket"] = old_bucket if old else "Not found"
        amend["new_bucket"] = new_bucket
        amend["old_base"] = old["base"] if old else None
        amend["old_igst"] = old["igst"] if old else None
        amend["old_cgst"] = old["cgst"] if old else None
        amend["old_sgst"] = old["sgst"] if old else None
        amend["old_cess"] = old["cess"] if old else None
        amend["revised_base"] = amend["base"]
        amend["delta"] = delta
        amend["bucket_delta"] = bucket_delta
        amend["signed_base"] = delta["base"]
        for comp in ("igst", "cgst", "sgst", "cess"):
            amend[f"signed_{comp}"] = delta[comp]
        amend["signed_gst"] = delta["igst"] + delta["cgst"] + delta["sgst"] + delta["cess"]
        amend["available_effect"] = bucket_delta["Available"]
        amend["rcm_effect"] = bucket_delta["RCM available"]
        amend["not_available_effect"] = bucket_delta["Not available"]
        amend["rcm_not_available_effect"] = bucket_delta["RCM not available"]
        amend["rejected_effect"] = bucket_delta["Rejected"]
        amend["unknown_effect"] = bucket_delta["Unknown"]
        amend["bucket"] = new_bucket
        amend["linked_original_id"] = old["id"] if old else ""
        amend["event_note"] = "" if old else "Original record not found in detail; revised amount treated as a full addition"
        if ambiguous:
            amend["event_note"] = "Multiple original records matched the amendment key; review the link"
        normalized.append(amend)
    return normalized


def attach_amendment_book_links(amendments, original_matches, purchase_records, note_records):
    by_id = {r["id"]: r for r in purchase_records + note_records}
    for a in amendments:
        old = a.get("original")
        book = by_id.get(old.get("book_id")) if old and old.get("book_id") else None
        if book is None:
            # If the original is missing/unmatched, try the revised document
            # reference against an unassigned book row. This is still a 2B
            # amendment record and is explicitly identified as such.
            pool = purchase_records if a["source_sheet"] == "B2BA" else note_records
            primary = "voucher"
            alternatives = ["invoice"] if a["source_sheet"] == "B2BA" else ["ref"]
            gst = a["gstin"]
            tok = norm_doc(a["doc_no"])
            cands = [b for b in pool if b["gstin"] == gst and tok and tok in [norm_doc(b.get(primary)), *[norm_doc(b.get(x)) for x in alternatives]]]
            if len(cands) == 1:
                book = cands[0]
        a["book_id"] = book["id"] if book else ""
        a["book_record"] = book
        a["book_month"] = book["month"] if book else ""
        if book:
            book.setdefault("amendments", []).append(a)


def read_source_summary(wb, sheet_name, row_num_1based):
    ws = wb.sheet_by_name(sheet_name)
    row = ws.row_values(row_num_1based - 1)
    return {"igst": number(row[3]), "cgst": number(row[4]), "sgst": number(row[5]), "cess": number(row[6]), "label": text(row[1])}


def build_source_controls(wb, b2b, notes, amendments, rejected):
    avail = {
        "b2b_nonrcm": read_source_summary(wb, "ITC Available", 10),
        "b2ba_nonrcm": read_source_summary(wb, "ITC Available", 13),
        "b2b_rcm": read_source_summary(wb, "ITC Available", 20),
        "cdnr_nonrcm": read_source_summary(wb, "ITC Available", 31),
        "cdnr_rcm": read_source_summary(wb, "ITC Available", 33),
    }
    not_avail = {
        "b2b_nonrcm": read_source_summary(wb, "ITC not available", 10),
        "cdnr_nonrcm": read_source_summary(wb, "ITC not available", 26),
        "cdnra_nonrcm": read_source_summary(wb, "ITC not available", 27),
    }
    rej_summary = read_source_summary(wb, "ITC Rejected", 10)

    def sums(records, predicate, fields=("igst", "cgst", "sgst", "cess")):
        chosen = [r for r in records if predicate(r)]
        return len(chosen), {f: sum(number(r.get(f, 0)) for r in chosen) for f in fields}, sum(number(r.get("base", 0)) for r in chosen)

    raw = {}
    raw["b2b_nonrcm"] = sums(b2b, lambda r: r["availability"].lower() == "yes" and r["rcm"] != "Y")
    raw["b2ba_nonrcm"] = sums(amendments, lambda r: r["source_sheet"] == "B2BA" and r["availability"].lower() == "yes" and r["rcm"] != "Y")
    raw["b2b_rcm"] = sums(b2b, lambda r: r["availability"].lower() == "yes" and r["rcm"] == "Y")
    raw["cdnr_nonrcm"] = sums(notes, lambda r: r["availability"].lower() == "yes" and r["rcm"] != "Y")
    raw["cdnr_rcm"] = sums(notes, lambda r: r["availability"].lower() == "yes" and r["rcm"] == "Y")
    raw["b2b_no"] = sums(b2b, lambda r: r["availability"].lower() == "no" and r["rcm"] != "Y")
    raw["cdnr_no"] = sums(notes, lambda r: r["availability"].lower() == "no" and r["rcm"] != "Y")
    raw["cdnra_no"] = sums(amendments, lambda r: r["source_sheet"] == "B2B-CDNRA" and r["availability"].lower() == "no" and r["rcm"] != "Y")
    raw["rejected"] = sums(rejected, lambda r: True)

    checks = [
        ("Available B2B invoices (non-RCM)", "B2B rows: Yes, RCM≠Y", raw["b2b_nonrcm"], avail["b2b_nonrcm"], "ITC Available!J? row 10"),
        ("Available B2B invoice amendments (revised values)", "B2BA revised detail: Yes, RCM≠Y", raw["b2ba_nonrcm"], avail["b2ba_nonrcm"], "ITC Available row 13; source summary shows the revised gross value"),
        ("Available B2B invoices (RCM)", "B2B rows: Yes, RCM=Y", raw["b2b_rcm"], avail["b2b_rcm"], "ITC Available row 20"),
        ("Available B2B credit notes (non-RCM)", "B2B-CDNR rows: Yes, RCM≠Y", raw["cdnr_nonrcm"], avail["cdnr_nonrcm"], "ITC Available row 31; raw C-note effect is signed negative in this workbook"),
        ("Available B2B credit notes (RCM)", "B2B-CDNR rows: Yes, RCM=Y", raw["cdnr_rcm"], avail["cdnr_rcm"], "ITC Available row 33; review because raw detail and summary differ"),
        ("Not-available B2B invoices", "B2B rows: No, RCM≠Y", raw["b2b_no"], not_avail["b2b_nonrcm"], "ITC not available row 10"),
        ("Not-available B2B credit notes", "B2B-CDNR rows: No, RCM≠Y", raw["cdnr_no"], not_avail["cdnr_nonrcm"], "ITC not available row 26; review if source summary excludes the raw note"),
        ("Not-available credit-note amendments (revised values)", "B2B-CDNRA revised detail: No, RCM≠Y", raw["cdnra_no"], not_avail["cdnra_nonrcm"], "ITC not available row 27; detail amendment is not added as a full amount"),
        ("Rejected B2B invoices", "B2B(Rejected) detail", raw["rejected"], rej_summary, "ITC Rejected row 10"),
    ]
    rows = []
    for label, detail_desc, detail, summary, comment in checks:
        count, detail_tax, base = detail
        differences = {k: detail_tax[k] - summary[k] for k in ("igst", "cgst", "sgst", "cess")}
        max_diff = max(abs(v) for v in differences.values())
        status = "OK" if max_diff <= TOLERANCE else "REVIEW"
        rows.append({
            "control": label, "detail_source": detail_desc, "raw_count": count,
            "raw_base": base, "raw_igst": detail_tax["igst"], "raw_cgst": detail_tax["cgst"],
            "raw_sgst": detail_tax["sgst"], "raw_cess": detail_tax["cess"],
            "summary_igst": summary["igst"], "summary_cgst": summary["cgst"],
            "summary_sgst": summary["sgst"], "summary_cess": summary["cess"],
            "diff_igst": differences["igst"], "diff_cgst": differences["cgst"],
            "diff_sgst": differences["sgst"], "diff_cess": differences["cess"],
            "status": status, "comment": comment,
        })
    return rows


def event_bucket_effect(event, bucket):
    if event.get("event") == "Amendment adjustment":
        return event.get("bucket_delta", {}).get(bucket, {k: 0.0 for k in ("base", "igst", "cgst", "sgst", "cess")})
    if event.get("bucket") == bucket:
        return {"base": event.get("signed_base", 0.0), "igst": event.get("signed_igst", 0.0),
                "cgst": event.get("signed_cgst", 0.0), "sgst": event.get("signed_sgst", 0.0),
                "cess": event.get("signed_cess", 0.0)}
    return {k: 0.0 for k in ("base", "igst", "cgst", "sgst", "cess")}


def raw_amendment_summary(amendment):
    return {"base": amendment["base"], "igst": amendment["igst"], "cgst": amendment["cgst"], "sgst": amendment["sgst"], "cess": amendment["cess"]}


def build_monthly_summary(purchases, book_notes, b2b, notes, amendments, rejected):
    months = {month_key(m): {
        "month": m,
        "book_purchase_taxed_base": 0.0, "book_purchase_no_tax_base": 0.0,
        "book_note_taxed_base": 0.0, "book_note_no_tax_base": 0.0,
        "book_purchase_cgst": 0.0, "book_purchase_sgst": 0.0, "book_purchase_igst": 0.0,
        "book_note_cgst": 0.0, "book_note_sgst": 0.0, "book_note_igst": 0.0,
        "book_purchase_count": 0, "book_note_count": 0,
        "book_no_tax_count": 0, "book_unmatched_count": 0, "book_unmatched_gst": 0.0,
        "b2b_matches_same": 0, "b2b_matches_late": 0, "b2b_matches_early": 0,
        "b2b_open_available_count": 0, "b2b_open_available_gst": 0.0,
        "b2b_open_rcm_count": 0, "b2b_open_not_available_count": 0,
        "b2b_open_rejected_count": 0,
    } for m in FY_MONTHS}
    for r in purchases:
        m = months.get(r["month"])
        if not m:
            continue
        m["book_purchase_count"] += 1
        if r["no_tax"]:
            m["book_purchase_no_tax_base"] += r["base"]
            m["book_no_tax_count"] += 1
        else:
            m["book_purchase_taxed_base"] += r["base"]
        m["book_purchase_cgst"] += r["cgst"]
        m["book_purchase_sgst"] += r["sgst"]
        m["book_purchase_igst"] += r["igst"]
        if not r.get("match_info") and not r.get("rejected_match"):
            m["book_unmatched_count"] += 1
            m["book_unmatched_gst"] += r["signed_gst"]
    for r in book_notes:
        m = months.get(r["month"])
        if not m:
            continue
        m["book_note_count"] += 1
        if r["no_tax"]:
            m["book_note_no_tax_base"] += r["no_tax_base_signed"]
            m["book_no_tax_count"] += 1
        else:
            m["book_note_taxed_base"] += r["taxed_base_signed"]
        m["book_note_cgst"] += r["signed_cgst"]
        m["book_note_sgst"] += r["signed_sgst"]
        m["book_note_igst"] += r["signed_igst"]
        if not r.get("match_info"):
            m["book_unmatched_count"] += 1
            m["book_unmatched_gst"] += r["signed_gst"]

    # Original GSTR-2B lines and amendments are grouped by the supplier's
    # reporting period. The amendment event uses revised minus original value.
    all_events = b2b + notes + amendments + rejected
    for e in all_events:
        m = months.get(month_key(e.get("period")))
        if not m:
            continue
        for bucket, prefix in (("Available", "b2b_available"), ("RCM available", "b2b_rcm"),
                               ("Not available", "b2b_not_available"), ("RCM not available", "b2b_rcm_not_available"),
                               ("Rejected", "b2b_rejected"), ("Unknown", "b2b_unknown")):
            eff = event_bucket_effect(e, bucket)
            for k in ("base", "igst", "cgst", "sgst", "cess"):
                m[f"{prefix}_{k}"] = m.get(f"{prefix}_{k}", 0.0) + eff[k]
        if e.get("event") != "Rejected" and e.get("source_sheet") in ("B2B", "B2B-CDNR"):
            book_id = e.get("book_id")
            if book_id:
                book = e.get("book_record")
                if book:
                    bmonth = book.get("month", "")
                    eperiod = month_key(e.get("period"))
                    if bmonth == eperiod:
                        m["b2b_matches_same"] += 1
                    elif bmonth < eperiod:
                        m["b2b_matches_late"] += 1
                    elif bmonth > eperiod:
                        m["b2b_matches_early"] += 1
            else:
                bucket = e.get("bucket")
                if bucket == "Available":
                    m["b2b_open_available_count"] += 1
                    m["b2b_open_available_gst"] += e.get("signed_gst", 0.0)
                elif bucket == "RCM available":
                    m["b2b_open_rcm_count"] += 1
                elif bucket in ("Not available", "RCM not available"):
                    m["b2b_open_not_available_count"] += 1
        elif e.get("event") == "Amendment adjustment" and not e.get("book_id"):
            if e.get("bucket") == "Available":
                m["b2b_open_available_count"] += 1
                m["b2b_open_available_gst"] += e.get("available_effect", {}).get("igst", 0) + e.get("available_effect", {}).get("cgst", 0) + e.get("available_effect", {}).get("sgst", 0) + e.get("available_effect", {}).get("cess", 0)
        elif e.get("event") == "Rejected" and not e.get("book_id"):
            m["b2b_open_rejected_count"] += 1

    rows = []
    for k, m in months.items():
        purchase_tax = m["book_purchase_cgst"] + m["book_purchase_sgst"] + m["book_purchase_igst"]
        note_tax = m["book_note_cgst"] + m["book_note_sgst"] + m["book_note_igst"]
        book_net_base = m["book_purchase_taxed_base"] + m["book_purchase_no_tax_base"] + m["book_note_taxed_base"] + m["book_note_no_tax_base"]
        book_taxed_base = m["book_purchase_taxed_base"] + m["book_note_taxed_base"]
        b2b_base = m.get("b2b_available_base", 0.0)
        b2b_igst = m.get("b2b_available_igst", 0.0)
        b2b_cgst = m.get("b2b_available_cgst", 0.0)
        b2b_sgst = m.get("b2b_available_sgst", 0.0)
        b2b_cess = m.get("b2b_available_cess", 0.0)
        b2b_gst = b2b_igst + b2b_cgst + b2b_sgst + b2b_cess
        book_igst = m["book_purchase_igst"] + m["book_note_igst"]
        book_cgst = m["book_purchase_cgst"] + m["book_note_cgst"]
        book_sgst = m["book_purchase_sgst"] + m["book_note_sgst"]
        book_gst = book_igst + book_cgst + book_sgst
        rcm_gst = sum(m.get(f"b2b_rcm_{c}", 0.0) for c in ("igst", "cgst", "sgst", "cess"))
        na_gst = sum(m.get(f"b2b_not_available_{c}", 0.0) for c in ("igst", "cgst", "sgst", "cess"))
        rejected_gst = sum(m.get(f"b2b_rejected_{c}", 0.0) for c in ("igst", "cgst", "sgst", "cess"))
        rows.append({
            "month": m["month"],
            "book_purchase_taxed_base": m["book_purchase_taxed_base"],
            "book_purchase_no_tax_base": m["book_purchase_no_tax_base"],
            "book_note_taxed_base": m["book_note_taxed_base"],
            "book_note_no_tax_base": m["book_note_no_tax_base"],
            "book_net_base": book_net_base,
            "book_taxed_base_net": book_taxed_base,
            "book_purchase_cgst": m["book_purchase_cgst"],
            "book_purchase_sgst": m["book_purchase_sgst"],
            "book_purchase_igst": m["book_purchase_igst"],
            "book_purchase_gst": purchase_tax,
            "book_note_cgst_signed": m["book_note_cgst"],
            "book_note_sgst_signed": m["book_note_sgst"],
            "book_note_igst_signed": m["book_note_igst"],
            "book_note_gst_signed": note_tax,
            "book_net_cgst": book_cgst, "book_net_sgst": book_sgst,
            "book_net_igst": book_igst, "book_net_gst": book_gst,
            "b2b_available_base": b2b_base,
            "b2b_available_cgst": b2b_cgst, "b2b_available_sgst": b2b_sgst,
            "b2b_available_igst": b2b_igst, "b2b_available_cess": b2b_cess,
            "b2b_available_gst": b2b_gst,
            "diff_taxed_base": book_taxed_base - b2b_base,
            "diff_cgst": book_cgst - b2b_cgst,
            "diff_sgst": book_sgst - b2b_sgst,
            "diff_igst": book_igst - b2b_igst,
            "diff_gst": book_gst - b2b_gst,
            "b2b_rcm_base": m.get("b2b_rcm_base", 0.0), "b2b_rcm_gst": rcm_gst,
            "b2b_not_available_base": m.get("b2b_not_available_base", 0.0), "b2b_not_available_gst": na_gst,
            "b2b_rejected_base": m.get("b2b_rejected_base", 0.0), "b2b_rejected_gst": rejected_gst,
            "book_no_tax_base_net": m["book_purchase_no_tax_base"] + m["book_note_no_tax_base"],
            "book_no_tax_count": m["book_no_tax_count"],
            "matched_same_count": m["b2b_matches_same"],
            "matched_book_earlier_count": m["b2b_matches_late"],
            "matched_book_later_count": m["b2b_matches_early"],
            "unmatched_book_count": m["book_unmatched_count"],
            "unmatched_book_gst": m["book_unmatched_gst"],
            "unmatched_2b_available_count": m["b2b_open_available_count"],
            "unmatched_2b_available_gst": m["b2b_open_available_gst"],
            "unmatched_2b_rcm_count": m["b2b_open_rcm_count"],
            "unmatched_2b_not_available_count": m["b2b_open_not_available_count"],
            "unmatched_2b_rejected_count": m["b2b_open_rejected_count"],
        })
    return rows


def build_month_flow(b2b, notes, amendments):
    grouped = {}
    for e in b2b + notes:
        book = e.get("book_record")
        if not book:
            continue
        bmonth, pmonth = book.get("month", ""), month_key(e.get("period"))
        bucket = e.get("bucket", "Unknown")
        key = (e["transaction"], bucket, bmonth, pmonth, "Original")
        g = grouped.setdefault(key, {"count": 0, "book_base": 0.0, "book_igst": 0.0, "book_cgst": 0.0, "book_sgst": 0.0, "book_gst": 0.0,
                                     "b2b_base": 0.0, "b2b_igst": 0.0, "b2b_cgst": 0.0, "b2b_sgst": 0.0, "b2b_cess": 0.0, "b2b_gst": 0.0})
        sign = book["signed"]
        g["count"] += 1
        g["book_base"] += sign * book["base"]
        g["book_igst"] += sign * book["igst"]
        g["book_cgst"] += sign * book["cgst"]
        g["book_sgst"] += sign * book["sgst"]
        g["book_gst"] += sign * (book["igst"] + book["cgst"] + book["sgst"])
        g["b2b_base"] += e["signed_base"]
        g["b2b_igst"] += e["signed_igst"]
        g["b2b_cgst"] += e["signed_cgst"]
        g["b2b_sgst"] += e["signed_sgst"]
        g["b2b_cess"] += e["signed_cess"]
        g["b2b_gst"] += e["signed_gst"]
    for e in amendments:
        book = e.get("book_record")
        if not book:
            continue
        bmonth, pmonth = book.get("month", ""), month_key(e.get("period"))
        key = (e["transaction"], f"Amendment → {e.get('new_bucket', e.get('bucket',''))}", bmonth, pmonth, "Amendment adjustment")
        g = grouped.setdefault(key, {"count": 0, "book_base": 0.0, "book_igst": 0.0, "book_cgst": 0.0, "book_sgst": 0.0, "book_gst": 0.0,
                                     "b2b_base": 0.0, "b2b_igst": 0.0, "b2b_cgst": 0.0, "b2b_sgst": 0.0, "b2b_cess": 0.0, "b2b_gst": 0.0})
        g["count"] += 1
        # Books recorded the original document once; the amendment month is a
        # 2B-only change to that document, not a second books posting.
        g["b2b_base"] += e["signed_base"]
        g["b2b_igst"] += e["signed_igst"]
        g["b2b_cgst"] += e["signed_cgst"]
        g["b2b_sgst"] += e["signed_sgst"]
        g["b2b_cess"] += e["signed_cess"]
        g["b2b_gst"] += e["signed_gst"]
    rows = []
    for key, v in sorted(grouped.items(), key=lambda item: (item[0][3], item[0][2], item[0][0], item[0][4])):
        transaction, bucket, bmonth, pmonth, event = key
        if event == "Amendment adjustment":
            timing = "Amendment adjustment (book posting shown for original month)"
        elif bmonth == pmonth:
            timing = "Same month"
        elif bmonth < pmonth:
            timing = "Book posted earlier; 2B reflected later"
        else:
            timing = "2B reflected earlier; book posted later"
        rows.append({"transaction": transaction, "bucket": bucket, "book_month": bmonth,
                     "2b_period": pmonth, "event": event, "timing": timing, **v,
                     "gst_variance": v["book_gst"] - v["b2b_gst"]})
    return rows


def build_recon_rows(book_records, matches_by_book, rejected_matches_by_book, original_lookup, amendments_by_original):
    rows = []
    for b in book_records:
        info = matches_by_book.get(b["id"])
        g = info if info else None
        reject = rejected_matches_by_book.get(b["id"])
        amendment_list = amendments_by_original.get(g["id"], []) if g else []
        a = amendment_list[0] if amendment_list else None
        if g:
            if g.get("availability", "").lower() == "no":
                status = "Matched - ITC marked No"
            elif g.get("rcm") == "Y":
                status = "Matched - RCM (separate)"
            else:
                status = "Matched"
            method = g.get("match_method", "")
            if g.get("match_duplicate"):
                status = "Matched - duplicate reference; review"
            elif g.get("match_gstin_mismatch"):
                status = "Matched - unique ref/base; GSTIN differs - review"
            base_comp = b["signed_base"] - g["signed_base"]
            tax_comp = {
                "igst": b["signed_igst"] - g["signed_igst"],
                "cgst": b["signed_cgst"] - g["signed_cgst"],
                "sgst": b["signed_sgst"] - g["signed_sgst"],
            }
            two_base = g["base"]
            two_igst, two_cgst, two_sgst, two_cess = g["igst"], g["cgst"], g["sgst"], g["cess"]
            two_period, two_filing, two_doc, two_date = g["period"], g["filing_date"], g["doc_no"], g["doc_date"]
            avail, rcm = g["availability"], g["rcm"]
            match_status = status
            timing = "Same month" if b["month"] == month_key(two_period) else ("Book posted earlier; 2B later" if b["month"] < month_key(two_period) else "2B earlier; book posted later")
        else:
            two_base = two_igst = two_cgst = two_sgst = two_cess = None
            two_period = two_filing = two_date = None
            two_doc = avail = rcm = ""
            method = ""
            base_comp = None
            tax_comp = {"igst": None, "cgst": None, "sgst": None}
            timing = ""
            if reject:
                match_status = "Rejected in GSTR-2B (not eligible)"
            elif b.get("cross_gstin_review"):
                match_status = "Cross-GSTIN candidate is ambiguous or already matched elsewhere; review"
            elif b["possible_duplicate"]:
                match_status = "Unmatched duplicate reference - review"
            elif b["no_tax"]:
                match_status = "No separate GST posted"
            elif not b["gstin"]:
                match_status = "No GSTIN in books"
            elif not (norm_doc(b.get("voucher")) or norm_doc(b.get("invoice"))):
                match_status = "No document reference in books"
            else:
                match_status = "Not located in 2B - timing/key/vendor filing review"
        row = {
            "book_id": b["id"], "branch": b["branch"], "source_file": b["source_file"], "source_row": b["source_row"],
            "book_date": b["date"], "book_month": b["date"], "supplier": b["party"], "gstin": b["gstin"],
            "voucher_type": b["voucher_type"], "voucher_no": b["voucher"], "supplier_invoice_no": b["invoice"],
            "supplier_invoice_date": b["doc_date"], "book_gross": b["gross"], "book_base": b["base"],
            "book_ledger_sum": b["ledger_sum"], "book_base_qa_delta": b["base_qa_delta"],
            "book_base_method": b["base_method"], "purchase_accounts_source": b["purchase_accounts_source"],
            "purchase_account_adjustment": b["purchase_account_adjustment"],
            "book_no_tax_base": b["base"] if b["no_tax"] else 0.0,
            "book_cgst": b["cgst"], "book_sgst": b["sgst"], "book_igst": b["igst"], "book_gst": b["gst"],
            "book_all_items": b["all_items"], "book_round_off": b["round_off"],
            "match_status": match_status, "match_method": method,
            "2b_supplier_gstin": g["gstin"] if g else "",
            "gstin_match": "Yes" if g and b["gstin"] == g["gstin"] else ("No - unique ref/base fallback" if g else ""),
            "2b_document_no": two_doc,
            "2b_document_date": two_date, "2b_period": two_period, "2b_filing_date": two_filing,
            "2b_itc_availability": avail, "2b_rcm": rcm, "2b_raw_base": two_base,
            "2b_igst": two_igst, "2b_cgst": two_cgst, "2b_sgst": two_sgst, "2b_cess": two_cess,
            "base_variance": base_comp, "igst_variance": tax_comp["igst"],
            "cgst_variance": tax_comp["cgst"], "sgst_variance": tax_comp["sgst"],
            "gst_variance": (tax_comp["igst"] + tax_comp["cgst"] + tax_comp["sgst"]) if g else None,
            "amendment": "Yes" if a else "No", "amendment_period": a.get("period") if a else None,
            "amendment_base_delta": a.get("signed_base") if a else None,
            "amendment_gst_delta": a.get("signed_gst") if a else None,
            "timing": timing, "possible_duplicate": "Yes" if b["possible_duplicate"] else "No",
            "rejected_2b_doc_no": reject.get("doc_no", "") if reject else "",
        }
        rows.append(row)
    return rows


def build_credit_note_rows(book_records, matches_by_book, amendments_by_original):
    rows = []
    for b in book_records:
        g = matches_by_book.get(b["id"])
        amendment_list = amendments_by_original.get(g["id"], []) if g else []
        a = amendment_list[0] if amendment_list else None
        if g:
            s = g.get("raw_sign", 1)
            book_base_effect = b["signed_base"]
            two_base_effect = g["signed_base"]
            book_tax = {"igst": b["signed_igst"], "cgst": b["signed_cgst"], "sgst": b["signed_sgst"]}
            two_tax = {k: g[f"signed_{k}"] for k in ("igst", "cgst", "sgst")}
            status = "Matched"
            if g.get("availability", "").lower() == "no":
                status = "Matched - ITC marked No"
            elif g.get("rcm") == "Y":
                status = "Matched - RCM (separate)"
            if g.get("match_duplicate"):
                status = "Matched - duplicate reference; review"
            elif g.get("match_gstin_mismatch"):
                status = "Matched - unique ref/base; GSTIN differs - review"
            base_variance = book_base_effect - two_base_effect
            igst_diff = book_tax["igst"] - two_tax["igst"]
            cgst_diff = book_tax["cgst"] - two_tax["cgst"]
            sgst_diff = book_tax["sgst"] - two_tax["sgst"]
            timing = "Same month" if b["month"] == month_key(g.get("period")) else ("Book posted earlier; 2B later" if b["month"] < month_key(g.get("period")) else "2B earlier; book posted later")
            method = g.get("match_method", "")
            two_doc = g["doc_no"]
            two_type = g["note_type"]
            two_date, two_period, two_filing = g["doc_date"], g["period"], g["filing_date"]
            availability, rcm = g["availability"], g["rcm"]
            raw_base, raw_igst, raw_cgst, raw_sgst = g["base"], g["igst"], g["cgst"], g["sgst"]
            signed_base_2b, signed_gst_2b = two_base_effect, g["signed_gst"]
        else:
            if b.get("cross_gstin_review"):
                status = "Cross-GSTIN candidate is ambiguous or already matched elsewhere; review"
            elif b["possible_duplicate"]:
                status = "Unmatched duplicate reference - review"
            else:
                status = "No GSTIN in books" if not b["gstin"] else "Not located in 2B - confirm supplier note"
            base_variance = igst_diff = cgst_diff = sgst_diff = None
            timing = method = two_doc = two_type = availability = rcm = ""
            two_date = two_period = two_filing = None
            raw_base = raw_igst = raw_cgst = raw_sgst = signed_base_2b = signed_gst_2b = None
            book_base_effect = b["signed_base"]
        rows.append({
            "book_id": b["id"], "branch": b["branch"], "source_file": b["source_file"], "source_row": b["source_row"],
            "book_date": b["date"], "book_month": b["date"], "supplier": b["party"], "gstin": b["gstin"],
            "voucher_no": b["voucher"], "voucher_ref_no": b["invoice"], "note_date": b["doc_date"],
            "book_gross": b["gross"], "book_base_raw": b["base"], "book_signed_base": book_base_effect,
            "book_ledger_sum": b["ledger_sum"], "book_base_qa_delta": b["base_qa_delta"],
            "book_base_method": b["base_method"], "purchase_accounts_source": b["purchase_accounts_source"],
            "purchase_account_adjustment": b["purchase_account_adjustment"],
            "book_no_tax_base_signed": b["no_tax_base_signed"], "book_cgst_signed": b["signed_cgst"],
            "book_sgst_signed": b["signed_sgst"], "book_igst_signed": b["signed_igst"], "book_gst_signed": b["signed_gst"],
            "match_status": status, "match_method": method,
            "2b_supplier_gstin": g["gstin"] if g else "",
            "gstin_match": "Yes" if g and b["gstin"] == g["gstin"] else ("No - unique ref/base fallback" if g else ""),
            "2b_note_no": two_doc, "2b_note_type": two_type,
            "2b_note_date": two_date, "2b_period": two_period, "2b_filing_date": two_filing,
            "2b_itc_availability": availability, "2b_rcm": rcm, "2b_raw_base": raw_base,
            "2b_raw_igst": raw_igst, "2b_raw_cgst": raw_cgst, "2b_raw_sgst": raw_sgst,
            "2b_signed_base_effect": signed_base_2b, "2b_signed_gst_effect": signed_gst_2b,
            "base_variance": base_variance, "igst_variance": igst_diff, "cgst_variance": cgst_diff,
            "sgst_variance": sgst_diff, "gst_variance": (igst_diff + cgst_diff + sgst_diff) if g else None,
            "amendment": "Yes" if a else "No", "amendment_period": a.get("period") if a else None,
            "amendment_base_delta": a.get("signed_base") if a else None,
            "amendment_gst_delta": a.get("signed_gst") if a else None,
            "timing": timing, "possible_duplicate": "Yes" if b["possible_duplicate"] else "No",
            "sign_note": "Book Debit Note treated as supplier Credit Note / negative ITC effect; confirm unmatched entries.",
        })
    return rows


def source_book_recon_row(record):
    return {
        "document_class": record["kind"], "book_id": record["id"], "branch": record["branch"],
        "source_file": record["source_file"], "source_row": record["source_row"],
        "book_date": record["date"], "book_month": record["date"], "supplier": record["party"],
        "gstin": record["gstin"], "voucher_no": record["voucher"], "other_reference": record["invoice"],
        "book_base_signed": record["signed_base"], "book_no_tax_base_signed": record["no_tax_base_signed"],
        "book_cgst_signed": record["signed_cgst"], "book_sgst_signed": record["signed_sgst"],
        "book_igst_signed": record["signed_igst"], "book_gst_signed": record["signed_gst"],
        "match_status": record.get("match_info", {}).get("gstr", {}).get("availability", "") if record.get("match_info") else "",
        "possible_duplicate": "Yes" if record.get("possible_duplicate") else "No",
    }


def build_gstr_detail_rows(b2b, notes, amendments, rejected):
    rows = []
    for e in b2b + notes + amendments + rejected:
        book = e.get("book_record")
        buckets = {}
        for bucket in ("Available", "RCM available", "Not available", "RCM not available", "Rejected", "Unknown"):
            buckets[bucket] = event_bucket_effect(e, bucket)
        if e.get("event") == "Amendment adjustment":
            raw_base = e.get("old_base")
            revised_base = e.get("revised_base")
        else:
            raw_base = revised_base = None
        if e.get("event") == "Rejected":
            reconciliation_status = "Rejected - not eligible"
        elif book and e.get("match_gstin_mismatch"):
            reconciliation_status = "Matched via unique ref/base; GSTIN differs - review"
        elif book:
            reconciliation_status = "Matched to book"
        elif e.get("cross_gstin_review"):
            reconciliation_status = "Cross-GSTIN candidate is ambiguous or already matched elsewhere; review"
        else:
            reconciliation_status = "Not located in books"
        rows.append({
            "record_id": e["id"], "source_sheet": e["source_sheet"], "source_row": e["source_row"],
            "event": e["event"], "transaction": e["transaction"], "gstin": e["gstin"],
            "supplier": e["supplier"], "document_no": e["doc_no"], "document_type": e.get("doc_type", ""),
            "note_type": e.get("note_type", ""), "document_date": e.get("doc_date"),
            "2b_period": e.get("period"), "supplier_filing_date": e.get("filing_date"),
            "itc_availability": e.get("availability", ""), "rcm": e.get("rcm", ""),
            "place_of_supply": e.get("pos", ""), "reason": e.get("reason", ""),
            "invoice_or_note_value": e.get("invoice_value"), "raw_taxable_base_revised": e.get("base"),
            "previous_base": raw_base, "revised_base": revised_base,
            "signed_base_effect": e.get("signed_base", 0.0),
            "signed_igst_effect": e.get("signed_igst", 0.0), "signed_cgst_effect": e.get("signed_cgst", 0.0),
            "signed_sgst_effect": e.get("signed_sgst", 0.0), "signed_cess_effect": e.get("signed_cess", 0.0),
            "signed_gst_effect": e.get("signed_gst", 0.0),
            "available_gst_effect": sum(buckets["Available"][k] for k in ("igst", "cgst", "sgst", "cess")),
            "rcm_available_gst_effect": sum(buckets["RCM available"][k] for k in ("igst", "cgst", "sgst", "cess")),
            "not_available_gst_effect": sum(buckets["Not available"][k] for k in ("igst", "cgst", "sgst", "cess")),
            "rcm_not_available_gst_effect": sum(buckets["RCM not available"][k] for k in ("igst", "cgst", "sgst", "cess")),
            "rejected_gst_effect": sum(buckets["Rejected"][k] for k in ("igst", "cgst", "sgst", "cess")),
            "old_itc_bucket": e.get("old_bucket", ""), "new_itc_bucket": e.get("new_bucket", ""),
            "old_document_no": e.get("old_doc_no", ""), "linked_original_id": e.get("linked_original_id", ""),
            "linked_book_id": book.get("id", "") if book else "",
            "book_gstin": book.get("gstin", "") if book else "",
            "gstin_match": "Yes" if book and book.get("gstin") == e.get("gstin") else ("No - unique ref/base fallback" if book else ""),
            "book_date": book.get("date") if book else None,
            "book_month": book.get("date") if book else None,
            "book_match_method": e.get("match_method", ""),
            "reconciliation_status": reconciliation_status,
            "amendment_note": e.get("event_note", ""),
        })
    return rows


def build_book_base_qa(purchases, notes, threshold=1.0):
    rows = []
    for b in purchases + notes:
        if abs(b.get("base_qa_delta", 0.0)) <= threshold:
            continue
        if abs(b.get("purchase_account_adjustment", 0.0)) > 0.005:
            review = "Purchase Accounts adjustment pattern applied; verify this source ledger represents a contra/withholding item."
        else:
            review = "Ledger-column sum differs from gross-derived base by more than ₹1; inspect the source voucher/account postings."
        rows.append({
            "document_class": b["kind"], "book_id": b["id"], "branch": b["branch"],
            "source_file": b["source_file"], "source_row": b["source_row"],
            "book_date": b["date"], "supplier": b["party"], "gstin": b["gstin"],
            "voucher_no": b["voucher"], "other_reference": b["invoice"],
            "gross_total": b["gross"], "all_items_source": b["all_items"],
            "purchase_accounts_source": b["purchase_accounts_source"],
            "purchase_account_adjustment_applied": b["purchase_account_adjustment"],
            "round_off": b["round_off"], "book_taxable_base": b["base"], "book_ledger_columns_sum": b["ledger_sum"],
            "ledger_sum_minus_base": b["base_qa_delta"], "base_method": b["base_method"],
            "cgst": b["cgst"], "sgst": b["sgst"], "igst": b["igst"], "gst_total": b["gst"],
            "review": review,
        })
    rows.sort(key=lambda r: abs(r["ledger_sum_minus_base"]), reverse=True)
    return rows


def build_open_book_rows(purchases, notes):
    rows = []
    for b in purchases + notes:
        if b.get("match_info") or b.get("rejected_match"):
            continue
        status = ""
        if b.get("cross_gstin_review"):
            status = "Cross-GSTIN document/base candidate is ambiguous or already matched elsewhere; review before linking"
        elif b.get("possible_duplicate"):
            status = "Duplicate/reused reference; one-to-one match went to another book row"
        elif b["no_tax"]:
            status = "No separate GST in tax columns; could be non-GST, ineligible/embedded tax, or tax not separately posted"
        elif not b["gstin"]:
            status = "GSTIN missing in books"
        else:
            status = "Book item not located in 2B; check filing period, invoice number, and supplier return"
        rows.append({
            "document_class": b["kind"], "book_id": b["id"], "branch": b["branch"],
            "source_file": b["source_file"], "source_row": b["source_row"],
            "book_date": b["date"], "book_month": b["date"], "supplier": b["party"],
            "supplier_gstin": b["gstin"], "voucher_no": b["voucher"],
            "supplier_invoice_or_ref": b["invoice"], "document_date": b["doc_date"],
            "book_gross": b["gross"], "book_base_signed": b["signed_base"],
            "book_ledger_sum": b["ledger_sum"], "book_base_qa_delta": b["base_qa_delta"],
            "book_base_method": b["base_method"], "purchase_accounts_source": b["purchase_accounts_source"],
            "purchase_account_adjustment": b["purchase_account_adjustment"],
            "book_no_tax_base_signed": b["no_tax_base_signed"],
            "book_cgst_signed": b["signed_cgst"], "book_sgst_signed": b["signed_sgst"],
            "book_igst_signed": b["signed_igst"], "book_gst_signed": b["signed_gst"],
            "no_separate_gst": "Yes" if b["no_tax"] else "No",
            "status_for_review": status,
        })
    return rows


def build_open_gstr_rows(b2b, notes, amendments, rejected):
    rows = []
    for e in b2b + notes + amendments + rejected:
        book = e.get("book_record")
        if book:
            continue
        # An amendment is not a second unmatched source document when its
        # original was already linked to books; if the original was not found,
        # it remains open for review.
        rows.append({
            "source_sheet": e["source_sheet"], "source_row": e["source_row"],
            "event": e["event"], "transaction": e["transaction"], "supplier_gstin": e["gstin"],
            "supplier": e["supplier"], "document_no": e["doc_no"], "document_type": e.get("doc_type", ""),
            "note_type": e.get("note_type", ""), "document_date": e.get("doc_date"),
            "2b_period": e.get("period"), "supplier_filing_date": e.get("filing_date"),
            "itc_availability": e.get("availability", ""), "rcm": e.get("rcm", ""),
            "bucket": e.get("bucket", ""), "raw_base": e.get("base"),
            "signed_base_effect": e.get("signed_base", 0.0),
            "signed_igst_effect": e.get("signed_igst", 0.0),
            "signed_cgst_effect": e.get("signed_cgst", 0.0),
            "signed_sgst_effect": e.get("signed_sgst", 0.0),
            "signed_cess_effect": e.get("signed_cess", 0.0),
            "signed_gst_effect": e.get("signed_gst", 0.0),
            "linked_original_id": e.get("linked_original_id", ""),
            "review_note": e.get("event_note", "") or ("Potential cross-GSTIN reference/base candidate was ambiguous or already assigned; left unmatched for review." if e.get("cross_gstin_review") else ""),
        })
    return rows


def month_row_to_list(r):
    return [r.get(h) for h in MONTH_HEADERS_KEYS]


MONTH_HEADERS_KEYS = [
    "month", "book_purchase_taxed_base", "book_purchase_no_tax_base", "book_note_taxed_base", "book_note_no_tax_base",
    "book_net_base", "book_taxed_base_net", "book_purchase_cgst", "book_purchase_sgst", "book_purchase_igst",
    "book_purchase_gst", "book_note_cgst_signed", "book_note_sgst_signed", "book_note_igst_signed",
    "book_note_gst_signed", "book_net_cgst", "book_net_sgst", "book_net_igst", "book_net_gst",
    "b2b_available_base", "b2b_available_cgst", "b2b_available_sgst", "b2b_available_igst", "b2b_available_cess",
    "b2b_available_gst", "diff_taxed_base", "diff_cgst", "diff_sgst", "diff_igst", "diff_gst",
    "b2b_rcm_base", "b2b_rcm_gst", "b2b_not_available_base", "b2b_not_available_gst",
    "b2b_rejected_base", "b2b_rejected_gst", "book_no_tax_base_net", "book_no_tax_count",
    "matched_same_count", "matched_book_earlier_count", "matched_book_later_count",
    "unmatched_book_count", "unmatched_book_gst", "unmatched_2b_available_count",
    "unmatched_2b_available_gst", "unmatched_2b_rcm_count", "unmatched_2b_not_available_count", "unmatched_2b_rejected_count",
]

MONTH_HEADERS = [
    "Month", "Book purchase base (GST separately posted)", "Book purchase base (no GST separately posted)",
    "Book debit-note base signed (GST posted)", "Book debit-note base signed (no GST posted)",
    "Book net base (all book entries)", "Book net CGST", "Book net SGST", "Book net IGST", "Book net GST",
    "2B standard available base (signed; amendment-adjusted)", "2B available CGST", "2B available SGST", "2B available IGST",
    "2B standard available GST (signed)", "Book less 2B base variance (GST-posted book rows only)",
    "Book less 2B CGST variance", "Book less 2B SGST variance", "Book less 2B IGST variance", "Book less 2B GST variance",
    "2B RCM GST effect (separate; claim after payment)", "2B ITC-not-available GST effect (signed)", "2B rejected GST (not eligible)",
    "Book no-separate-GST entries (count)", "Matched docs in same month (count)",
    "Book month earlier than 2B period (count)", "2B period earlier than book month (count)",
    "Unmatched book entries (count)", "Unmatched book GST (signed)",
    "Unmatched 2B standard available docs (count)", "Unmatched 2B standard available GST (signed)",
    "Unmatched 2B RCM docs (count)", "Unmatched 2B not-available docs (count)", "Unmatched 2B rejected docs (count)",
]


def build_month_flow_rows(flow):
    return flow


def to_sheet_row(data, headers):
    return [data.get(h) for h in headers]


def style_formats(workbook):
    fmt = {}
    fmt["title"] = workbook.add_format({"bold": True, "font_size": 19, "font_color": "#FFFFFF", "bg_color": "#17365D", "valign": "vcenter"})
    fmt["subtitle"] = workbook.add_format({"font_size": 10, "font_color": "#44546A", "italic": True, "text_wrap": True, "valign": "vcenter"})
    fmt["section"] = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#0F6B78", "font_size": 11, "valign": "vcenter"})
    fmt["card_label"] = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#4472C4", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"})
    fmt["card_value"] = workbook.add_format({"bold": True, "font_size": 16, "font_color": "#17365D", "bg_color": "#D9EAF7", "align": "center", "valign": "vcenter", "num_format": '#,##0.00;[Red](#,##0.00);-' , "border": 1, "border_color": "#FFFFFF"})
    fmt["card_count"] = workbook.add_format({"bold": True, "font_size": 16, "font_color": "#17365D", "bg_color": "#E2F0D9", "align": "center", "valign": "vcenter", "num_format": '#,##0;[Red](#,##0);-' , "border": 1, "border_color": "#FFFFFF"})
    fmt["table_header"] = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#17365D", "text_wrap": True, "valign": "vcenter", "border": 1})
    fmt["date"] = workbook.add_format({"num_format": "dd-mmm-yyyy"})
    fmt["month"] = workbook.add_format({"num_format": "mmm-yyyy"})
    fmt["money"] = workbook.add_format({"num_format": '#,##0.00;[Red](#,##0.00);-'})
    fmt["integer"] = workbook.add_format({"num_format": '#,##0;[Red](#,##0);-'})
    fmt["wrap"] = workbook.add_format({"text_wrap": True, "valign": "top"})
    fmt["note"] = workbook.add_format({"font_color": "#666666", "italic": True, "text_wrap": True, "valign": "top"})
    fmt["total_label"] = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "top": 1, "num_format": "mmm-yyyy"})
    fmt["total_money"] = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "top": 1, "num_format": '#,##0.00;[Red](#,##0.00);-'})
    fmt["total_int"] = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "top": 1, "num_format": '#,##0;[Red](#,##0);-'})
    fmt["positive"] = workbook.add_format({"font_color": "#9C0006", "bg_color": "#FFC7CE"})
    fmt["ok"] = workbook.add_format({"font_color": "#006100", "bg_color": "#C6EFCE"})
    fmt["review"] = workbook.add_format({"font_color": "#9C6500", "bg_color": "#FFEB9C"})
    fmt["body"] = workbook.add_format({"valign": "top"})
    return fmt


def choose_cell_format(key, value, formats):
    key_low = key.lower()
    if isinstance(value, datetime) or isinstance(value, date):
        return formats["month"] if "month" in key_low or "period" in key_low else formats["date"]
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if any(s in key_low for s in ("count", "row", "source row")):
            return formats["integer"]
        if key_low in ("year",):
            return formats["integer"]
        return formats["money"]
    return formats["wrap"] if len(text(value)) > 40 else None


def write_value(ws, row, col, value, key, formats):
    if value is None or value == "":
        ws.write_blank(row, col, None)
    elif isinstance(value, datetime):
        ws.write_datetime(row, col, value, choose_cell_format(key, value, formats))
    elif isinstance(value, date):
        ws.write_datetime(row, col, datetime(value.year, value.month, value.day), choose_cell_format(key, value, formats))
    elif isinstance(value, bool):
        ws.write_boolean(row, col, value)
    elif isinstance(value, (int, float)):
        ws.write_number(row, col, float(value), choose_cell_format(key, value, formats))
    else:
        ws.write_string(row, col, text(value), choose_cell_format(key, value, formats))


def width_for_header(header):
    low = header.lower()
    if any(x in low for x in ("source file", "supplier filing date", "match status", "review", "timing", "comment", "reason", "status", "method", "sign note", "availability")):
        return 27
    if any(x in low for x in ("supplier", "party", "document no", "invoice no", "voucher no", "gstin", "reference", "bucket", "event", "transaction", "class")):
        return 22
    if "month" in low or "period" in low or "date" in low:
        return 14
    if "count" in low or "row" in low:
        return 11
    if any(x in low for x in ("gst", "base", "gross", "amount", "variance", "igst", "cgst", "sgst", "cess", "value", "effect", "round")):
        return 16
    return 16


def add_table_sheet(workbook, name, title, subtitle, headers, rows, formats, freeze_cols=4, widths=None, tab_color="#4472C4"):
    ws = workbook.add_worksheet(name)
    ws.hide_gridlines(2)
    ws.set_tab_color(tab_color)
    last_col = max(0, len(headers) - 1)
    ws.merge_range(0, 0, 0, last_col, title, formats["title"])
    ws.set_row(0, 29)
    if subtitle:
        ws.merge_range(1, 0, 1, last_col, subtitle, formats["subtitle"])
        ws.set_row(1, 34)
    header_row = 3
    for idx, row in enumerate(rows, start=header_row + 1):
        for col, (header, value) in enumerate(zip(headers, row)):
            write_value(ws, idx, col, value, header, formats)
    last_row = header_row + len(rows)
    columns = [{"header": h, "header_format": formats["table_header"]} for h in headers]
    ws.add_table(header_row, 0, last_row, last_col, {"columns": columns, "style": "Table Style Medium 2", "autofilter": True})
    ws.freeze_panes(header_row + 1, freeze_cols)
    ws.set_row(header_row, 34)
    ws.set_default_row(18)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    ws.set_margins(0.25, 0.25, 0.5, 0.5)
    ws.set_footer("&LDrive Trucking | GSTR-2B vs Books&CPage &P of &N")
    if widths:
        for col, width in enumerate(widths):
            ws.set_column(col, col, width)
    else:
        for col, header in enumerate(headers):
            ws.set_column(col, col, width_for_header(header))
    for col, header in enumerate(headers):
        low = header.lower()
        if any(s in low for s in ("variance", "difference", "book less 2b")):
            if len(rows):
                ws.conditional_format(header_row + 1, col, last_row, col, {"type": "3_color_scale", "min_color": "#63BE7B", "mid_color": "#FFEB84", "max_color": "#F8696B"})
        elif "status" in low:
            if len(rows):
                ws.conditional_format(header_row + 1, col, last_row, col, {"type": "text", "criteria": "containing", "value": "review", "format": formats["review"]})
    return ws, header_row, last_row


def create_month_table_rows(months):
    output = []
    for r in months:
        output.append([
            r["month"], r["book_purchase_taxed_base"], r["book_purchase_no_tax_base"], r["book_note_taxed_base"],
            r["book_note_no_tax_base"], r["book_net_base"], r["book_net_cgst"], r["book_net_sgst"],
            r["book_net_igst"], r["book_net_gst"], r["b2b_available_base"], r["b2b_available_cgst"],
            r["b2b_available_sgst"], r["b2b_available_igst"], r["b2b_available_gst"], r["diff_taxed_base"],
            r["diff_cgst"], r["diff_sgst"], r["diff_igst"], r["diff_gst"], r["b2b_rcm_gst"],
            r["b2b_not_available_gst"], r["b2b_rejected_gst"], r["book_no_tax_count"],
            r["matched_same_count"], r["matched_book_earlier_count"], r["matched_book_later_count"],
            r["unmatched_book_count"], r["unmatched_book_gst"], r["unmatched_2b_available_count"],
            r["unmatched_2b_available_gst"], r["unmatched_2b_rcm_count"], r["unmatched_2b_not_available_count"],
            r["unmatched_2b_rejected_count"],
        ])
    total = ["FY Total"]
    for col in range(1, len(MONTH_HEADERS)):
        values = [row[col] for row in output]
        if col in (23, 24, 25, 26, 27, 29, 31, 32, 33):
            total.append(sum(number(v) for v in values))
        else:
            total.append(sum(number(v) for v in values))
    return output, total


def build_dashboard(workbook, months, summary, counts, control_rows, formats):
    ws = workbook.add_worksheet("Dashboard")
    ws.hide_gridlines(2)
    ws.set_tab_color("#17365D")
    ws.set_column("A:A", 20)
    ws.set_column("B:B", 17)
    ws.set_column("C:C", 4)
    ws.set_column("D:D", 18)
    ws.set_column("E:E", 17)
    ws.set_column("F:F", 4)
    ws.set_column("G:G", 18)
    ws.set_column("H:H", 17)
    ws.set_column("I:I", 4)
    ws.set_column("J:J", 18)
    ws.set_column("K:K", 17)
    ws.merge_range("A1:K1", "GSTR-2B vs Books Reconciliation", formats["title"])
    ws.set_row(0, 32)
    ws.merge_range("A2:K2", "Drive Trucking Private Limited | FY 2025-26 | Recipient GSTIN 24AAJCD4457C1ZV | Source GSTR-2B generated 29-Sep-2026", formats["subtitle"])
    ws.set_row(1, 24)

    kpis = [
        ("Book net GST (CGST+SGST+IGST)", summary["book_gst"]),
        ("2B standard available net GST (signed)", summary["eligible_gst"]),
        ("Book less 2B standard GST", summary["difference_gst"]),
        ("Book base with no separate GST posting", summary["no_tax_base"]),
        ("2B RCM GST effect (separate)", summary["rcm_gst"]),
    ]
    card_cols = [(0, 1), (2, 3), (4, 5), (6, 7), (8, 10)]
    for (label, value), (c0, c1) in zip(kpis, card_cols):
        ws.merge_range(3, c0, 3, c1, label, formats["card_label"])
        ws.merge_range(4, c0, 4, c1, value, formats["card_value"])
    ws.set_row(3, 35)
    ws.set_row(4, 30)

    chart = workbook.add_chart({"type": "column"})
    cat_range = "='Monthly Summary'!$A$5:$A$16"
    chart.add_series({"name": "Book net GST", "categories": cat_range,
                      "values": "='Monthly Summary'!$J$5:$J$16", "fill": {"color": "#4472C4"}, "border": {"none": True}})
    chart.add_series({"name": "2B standard available GST", "categories": cat_range,
                      "values": "='Monthly Summary'!$O$5:$O$16", "fill": {"color": "#70AD47"}, "border": {"none": True}})
    line = workbook.add_chart({"type": "line"})
    line.add_series({"name": "Book less 2B difference", "categories": cat_range,
                     "values": "='Monthly Summary'!$T$5:$T$16", "y2_axis": True,
                     "line": {"color": "#ED7D31", "width": 2.25}, "marker": {"type": "circle", "size": 4}})
    chart.combine(line)
    chart.set_title({"name": "Monthly net GST: books vs standard 2B available ITC"})
    chart.set_x_axis({"name": "Month", "date_axis": True, "num_format": "mmm-yy", "label_position": "low"})
    chart.set_y_axis({"name": "GST amount (₹)", "num_format": "#,##0;[Red](#,##0)", "major_gridlines": {"visible": True}})
    chart.set_y2_axis({"name": "Difference (₹)", "num_format": "#,##0;[Red](#,##0)"})
    chart.set_legend({"position": "bottom"})
    chart.set_size({"width": 880, "height": 400})
    chart.set_style(10)
    ws.insert_chart("A7", chart)

    ws.merge_range("A28:B28", "Reconciliation snapshot", formats["section"])
    snapshot_headers = ["Metric", "Count / amount"]
    snapshot_rows = [
        ("Purchase-register rows", counts["purchase_rows"]),
        ("Debit-note-register rows", counts["note_rows"]),
        ("Original GSTR-2B B2B invoices", counts["b2b_rows"]),
        ("Original GSTR-2B credit/debit notes", counts["note_2b_rows"]),
        ("B2B invoice amendments", counts["invoice_amendments"]),
        ("Credit/debit-note amendments", counts["note_amendments"]),
        ("IMS-rejected B2B records", counts["rejected_rows"]),
        ("Invoice pairs matched one-to-one", counts["matched_invoice_pairs"]),
        ("Credit-note pairs matched one-to-one", counts["matched_note_pairs"]),
        ("Unique ref/base matches with GSTIN mismatch (review)", counts["gstin_mismatch_matches"]),
        ("Open book entries", counts["open_book_rows"]),
        ("Book base QA rows with ledger-sum difference > ₹1", counts["base_qa_rows"]),
        ("2B original records not located in books", counts["open_2b_original_rows"]),
        ("Source-summary controls needing review", sum(1 for x in control_rows if x["status"] == "REVIEW")),
    ]
    ws.add_table(29, 0, 29 + len(snapshot_rows), 1, {
        "columns": [{"header": snapshot_headers[0]}, {"header": snapshot_headers[1]}],
        "data": snapshot_rows, "style": "Table Style Medium 2",
    })
    ws.set_column("A:A", 39)
    ws.set_column("B:B", 20)
    # add_table has been supplied its own data, so types/formatting are handled by XlsxWriter.
    ws.merge_range("D28:K28", "How to read this dashboard", formats["section"])
    notes = [
        "Positive book-less-2B difference = book GST ledger higher than standard non-RCM ITC in 2B. Negative = 2B higher.",
        "Book Debit Note Register amounts are signed negative as supplier credit-note reductions, consistent with the matching sample notes.",
        "2B credit notes (type C) are negative; supplier debit notes (type D) are positive. Amendments are revised less original, in the amendment period.",
        "Calculated book base uses Gross Total + All Items − GST − Round Off; a narrow Purchase Accounts contra/withholding adjustment is applied only for the identified twice-the-source-amount pattern. Ledger-sum differences are visible on Book Base QA.",
        "No-tax-base is kept separately: source books have no CGST/SGST/IGST split on those entries, so their calculated taxable/total amount is shown separately and not included in 2B GST totals.",
        "RCM, ITC-not-available, and rejected records are shown separately; see Source Control and Method & Notes before relying on any total.",
        "A cross-GSTIN match is allowed only for a unique normalized document reference plus exact-to-paise taxable base on both source sides; every such match is visibly flagged for GSTIN review.",
        "Use Month Flow to see documents booked in one month but reflected in another month's supplier GSTR-1 period.",
    ]
    for i, note in enumerate(notes, start=29):
        ws.merge_range(i, 3, i, 10, "• " + note, formats["note"])
        ws.set_row(i, 36)
    ws.freeze_panes(6, 0)
    ws.set_footer("&LDrive Trucking | GSTR-2B vs Books&CPage &P of &N")
    return ws


def write_workbook(months, invoice_rows, note_rows, gstr_rows, open_book, base_qa, open_gstr, flow_rows, control_rows, counts, summary):
    workbook = xlsxwriter.Workbook(str(OUTPUT), {"constant_memory": False})
    workbook.set_properties({
        "title": "GSTR-2B vs Books Reconciliation FY 2025-26",
        "subject": "Monthly purchase and supplier credit-note reconciliation",
        "author": "OpenAI / Arena.ai Agent Mode",
        "company": "Drive Trucking Private Limited",
        "comments": "Built from the uploaded books and GSTR-2B source workbooks. See Method & Notes.",
    })
    formats = style_formats(workbook)
    # Add the dashboard first so it opens as the first worksheet. Chart links may
    # point to the monthly sheet that is added immediately below.
    dashboard_ws = build_dashboard(workbook, months, summary, counts, control_rows, formats)

    # Summary rows and annual totals.
    month_data, month_total = create_month_table_rows(months)
    monthly_ws, monthly_header_row, monthly_last_row = add_table_sheet(
        workbook, "Monthly Summary", "Monthly reconciliation",
        "Book month = register Date. 2B month = supplier GSTR-1/IFF/1A/GSTR-5 Period in the uploaded quarterly 2B file. Credit notes are signed negative; amendments use revised minus original.",
        MONTH_HEADERS, month_data, formats, freeze_cols=1, tab_color="#0F6B78")
    total_row = monthly_last_row + 1
    for col, value in enumerate(month_total):
        write_value(monthly_ws, total_row, col, value, MONTH_HEADERS[col] if col < len(MONTH_HEADERS) else "", formats)
        if col == 0:
            monthly_ws.write_string(total_row, col, "FY Total", formats["total_label"])
        elif isinstance(value, (int, float)):
            is_count = any(k in MONTH_HEADERS[col].lower() for k in ("count", "entries", "docs"))
            monthly_ws.write_number(total_row, col, number(value), formats["total_int"] if is_count else formats["total_money"])
    monthly_ws.set_row(total_row, 22)
    monthly_ws.write(total_row + 2, 0, "Reconciliation note:", formats["section"])
    monthly_ws.merge_range(total_row + 2, 1, total_row + 2, len(MONTH_HEADERS) - 1,
        "The primary GST variance compares book GST ledger postings (CGST/SGST/IGST, including signed book debit notes) with standard non-RCM GSTR-2B available effects. Book base uses the gross-derived method documented on Method & Notes; ledger-sum QA exceptions are on Book Base QA. RCM, ITC-not-available, rejected, no-separate-GST base, timing flows, and unmatched records are disclosed separately.", formats["note"])
    for col, header in enumerate(MONTH_HEADERS):
        monthly_ws.set_column(col, col, width_for_header(header))
    monthly_ws.set_column(0, 0, 14)
    monthly_ws.freeze_panes(monthly_header_row + 1, 1)

    invoice_headers = [
        "Book ID", "Branch", "Source File", "Source Row", "Book Date", "Book Month", "Supplier", "Supplier GSTIN",
        "Voucher Type", "Voucher No.", "Supplier Invoice No.", "Supplier Invoice Date", "Book Gross Total",
        "Book Taxable/Base Amount (calculated)", "Book Ledger Columns Sum", "Ledger Sum − Base QA Difference", "Base Calculation Method",
        "Purchase Accounts Source", "Purchase Accounts Adjustment Applied", "Book No-Tax Base", "Book CGST", "Book SGST", "Book IGST", "Book GST Total",
        "All Items / Other (source field)", "Round Off", "Reconciliation Status", "Match Method", "2B Supplier GSTIN", "GSTIN Match?",
        "2B Invoice No.", "2B Invoice Date", "2B Period", "Supplier Filing Date", "2B ITC Availability", "2B RCM", "2B Raw Taxable Base",
        "2B IGST", "2B CGST", "2B SGST", "2B Cess", "Book less 2B Base Variance", "Book less 2B IGST Variance",
        "Book less 2B CGST Variance", "Book less 2B SGST Variance", "Book less 2B GST Variance", "2B Amendment?",
        "Amendment Period", "Amendment Base Delta", "Amendment GST Delta", "Timing Status", "Possible Duplicate Ref",
        "Rejected 2B Invoice No.",
    ]
    invoice_rows_values = [to_sheet_row(r, [
        "book_id", "branch", "source_file", "source_row", "book_date", "book_month", "supplier", "gstin", "voucher_type", "voucher_no",
        "supplier_invoice_no", "supplier_invoice_date", "book_gross", "book_base", "book_ledger_sum", "book_base_qa_delta", "book_base_method",
        "purchase_accounts_source", "purchase_account_adjustment", "book_no_tax_base", "book_cgst", "book_sgst", "book_igst", "book_gst",
        "book_all_items", "book_round_off", "match_status", "match_method", "2b_supplier_gstin", "gstin_match", "2b_document_no", "2b_document_date", "2b_period", "2b_filing_date",
        "2b_itc_availability", "2b_rcm", "2b_raw_base", "2b_igst", "2b_cgst", "2b_sgst", "2b_cess", "base_variance", "igst_variance", "cgst_variance",
        "sgst_variance", "gst_variance", "amendment", "amendment_period", "amendment_base_delta", "amendment_gst_delta", "timing", "possible_duplicate", "rejected_2b_doc_no",
    ]) for r in invoice_rows]
    add_table_sheet(workbook, "Book Invoice Recon", "Book purchase register vs GSTR-2B invoices",
                    "All purchase-register rows are included. Filter Reconciliation Status / GSTIN Match? / Timing Status to review exceptions. Unique cross-GSTIN reference-and-base matches are linked but explicitly flagged for supplier-registration review. Calculated book base = Gross Total + All Items − GST − Round Off, with a narrowly-triggered Purchase Accounts adjustment for the identified contra/withholding pattern. Ledger-column sum and its difference from the calculated base are retained for QA; see Book Base QA.",
                    invoice_headers, invoice_rows_values, formats, freeze_cols=6, tab_color="#4472C4")

    note_headers = [
        "Book ID", "Branch", "Source File", "Source Row", "Book Date", "Book Month", "Supplier", "Supplier GSTIN",
        "Voucher No.", "Voucher Ref No.", "Note Date", "Book Gross Total", "Book Base Raw", "Book Signed Base",
        "Book Ledger Columns Sum", "Ledger Sum − Base QA Difference", "Base Calculation Method", "Purchase Accounts Source",
        "Purchase Accounts Adjustment Applied", "Book No-Tax Base Signed", "Book CGST Signed", "Book SGST Signed", "Book IGST Signed", "Book GST Signed",
        "Reconciliation Status", "Match Method", "2B Supplier GSTIN", "GSTIN Match?", "2B Note No.", "2B Note Type", "2B Note Date", "2B Period",
        "Supplier Filing Date", "2B ITC Availability", "2B RCM", "2B Raw Base", "2B Raw IGST", "2B Raw CGST",
        "2B Raw SGST", "2B Signed Base Effect", "2B Signed GST Effect", "Book less 2B Base Variance",
        "Book less 2B IGST Variance", "Book less 2B CGST Variance", "Book less 2B SGST Variance", "Book less 2B GST Variance",
        "2B Amendment?", "Amendment Period", "Amendment Base Delta", "Amendment GST Delta", "Timing Status",
        "Possible Duplicate Ref", "Sign Convention Note",
    ]
    note_rows_values = [to_sheet_row(r, [
        "book_id", "branch", "source_file", "source_row", "book_date", "book_month", "supplier", "gstin", "voucher_no", "voucher_ref_no",
        "note_date", "book_gross", "book_base_raw", "book_signed_base", "book_ledger_sum", "book_base_qa_delta", "book_base_method",
        "purchase_accounts_source", "purchase_account_adjustment", "book_no_tax_base_signed", "book_cgst_signed", "book_sgst_signed",
        "book_igst_signed", "book_gst_signed", "match_status", "match_method", "2b_supplier_gstin", "gstin_match", "2b_note_no", "2b_note_type", "2b_note_date", "2b_period",
        "2b_filing_date", "2b_itc_availability", "2b_rcm", "2b_raw_base", "2b_raw_igst", "2b_raw_cgst", "2b_raw_sgst",
        "2b_signed_base_effect", "2b_signed_gst_effect", "base_variance", "igst_variance", "cgst_variance", "sgst_variance", "gst_variance",
        "amendment", "amendment_period", "amendment_base_delta", "amendment_gst_delta", "timing", "possible_duplicate", "sign_note",
    ]) for r in note_rows]
    add_table_sheet(workbook, "Book Credit Note Recon", "Book debit-note register vs GSTR-2B supplier notes",
                    "The uploaded books contain Debit Note registers. These are signed negative for purchases/ITC based on the matching GSTR-2B supplier credit-note examples. Type C in 2B is negative; type D is positive. Unique cross-GSTIN note matches based on reference and paise-exact taxable base are flagged for review. Book base uses the gross-derived method; ledger-sum differences are exposed for QA. Review unmatched and duplicate note references.",
                    note_headers, note_rows_values, formats, freeze_cols=6, tab_color="#C65911")

    gstr_headers = [
        "Record ID", "Source Sheet", "Source Row", "Event", "Transaction", "Supplier GSTIN", "Supplier",
        "Document No.", "Document Type", "Note Type", "Document Date", "2B Period", "Supplier Filing Date",
        "ITC Availability", "RCM", "Place of Supply", "Reason", "Invoice/Note Value", "Raw Revised Taxable Base",
        "Previous Base (amendment)", "Revised Base (amendment)", "Signed Base Effect", "Signed IGST Effect",
        "Signed CGST Effect", "Signed SGST Effect", "Signed Cess Effect", "Signed GST Effect",
        "Standard Available GST Effect", "RCM Available GST Effect", "Not-Available GST Effect",
        "RCM Not-Available GST Effect", "Rejected GST Effect", "Old ITC Bucket", "New ITC Bucket",
        "Old Document No.", "Linked Original ID", "Linked Book ID", "Book GSTIN", "GSTIN Match?", "Book Date", "Book Month",
        "Book Match Method", "Reconciliation Status", "Amendment Note",
    ]
    gstr_rows_values = [to_sheet_row(r, [
        "record_id", "source_sheet", "source_row", "event", "transaction", "gstin", "supplier", "document_no", "document_type", "note_type",
        "document_date", "2b_period", "supplier_filing_date", "itc_availability", "rcm", "place_of_supply", "reason", "invoice_or_note_value",
        "raw_taxable_base_revised", "previous_base", "revised_base", "signed_base_effect", "signed_igst_effect", "signed_cgst_effect", "signed_sgst_effect",
        "signed_cess_effect", "signed_gst_effect", "available_gst_effect", "rcm_available_gst_effect", "not_available_gst_effect",
        "rcm_not_available_gst_effect", "rejected_gst_effect", "old_itc_bucket", "new_itc_bucket", "old_document_no", "linked_original_id",
        "linked_book_id", "book_gstin", "gstin_match", "book_date", "book_month", "book_match_method", "reconciliation_status", "amendment_note",
    ]) for r in gstr_rows]
    add_table_sheet(workbook, "GSTR2B Detail", "GSTR-2B transaction detail (normalized for monthly comparison)",
                    "Original invoices/notes are kept in their supplier reporting period. B2BA/CDNRA rows are shown as revised-minus-original adjustments in their amended period; original and revised values remain visible. Unique reference/base matches across different supplier GSTINs are linked but explicitly flagged for review. Credit note type C is negative.",
                    gstr_headers, gstr_rows_values, formats, freeze_cols=6, tab_color="#70AD47")

    open_book_headers = [
        "Document Class", "Book ID", "Branch", "Source File", "Source Row", "Book Date", "Book Month", "Supplier", "Supplier GSTIN",
        "Voucher No.", "Supplier Invoice / Note Ref", "Document Date", "Book Gross", "Book Base Signed", "Book Ledger Columns Sum",
        "Ledger Sum − Base QA Difference", "Base Calculation Method", "Purchase Accounts Source", "Purchase Accounts Adjustment Applied",
        "Book No-Tax Base Signed", "Book CGST Signed", "Book SGST Signed", "Book IGST Signed", "Book GST Signed", "No Separate GST?", "Status for Review",
    ]
    open_book_rows_values = [to_sheet_row(r, [
        "document_class", "book_id", "branch", "source_file", "source_row", "book_date", "book_month", "supplier", "supplier_gstin",
        "voucher_no", "supplier_invoice_or_ref", "document_date", "book_gross", "book_base_signed", "book_ledger_sum", "book_base_qa_delta",
        "book_base_method", "purchase_accounts_source", "purchase_account_adjustment", "book_no_tax_base_signed",
        "book_cgst_signed", "book_sgst_signed", "book_igst_signed", "book_gst_signed", "no_separate_gst", "status_for_review",
    ]) for r in open_book]
    add_table_sheet(workbook, "Open Book Items", "Book rows not matched to an original 2B record",
                    "Includes no-tax-ledger items, GSTIN/reference gaps, duplicate references not assigned twice, and book GST not located in the provided 2B data. This is a review list, not a conclusion that ITC is unavailable.",
                    open_book_headers, open_book_rows_values, formats, freeze_cols=6, tab_color="#ED7D31")

    base_qa_headers = [
        "Document Class", "Book ID", "Branch", "Source File", "Source Row", "Book Date", "Supplier", "GSTIN",
        "Voucher No.", "Other Reference", "Gross Total", "All Items Source", "Purchase Accounts Source",
        "Purchase Accounts Adjustment Applied", "Round Off", "Calculated Book Taxable/Base", "Ledger Columns Sum",
        "Ledger Sum − Base QA Difference", "Base Calculation Method", "CGST", "SGST", "IGST", "GST Total", "Review",
    ]
    base_qa_rows_values = [to_sheet_row(r, [
        "document_class", "book_id", "branch", "source_file", "source_row", "book_date", "supplier", "gstin",
        "voucher_no", "other_reference", "gross_total", "all_items_source", "purchase_accounts_source",
        "purchase_account_adjustment_applied", "round_off", "book_taxable_base", "book_ledger_columns_sum",
        "ledger_sum_minus_base", "base_method", "cgst", "sgst", "igst", "gst_total", "review",
    ]) for r in base_qa]
    add_table_sheet(workbook, "Book Base QA", "Ledger-sum vs calculated book-base review",
                    "Rows with an absolute ledger-column-sum difference over ₹1. Calculated base = Gross Total + All Items − GST − Round Off, with the narrowly identified Purchase Accounts contra/withholding adjustment where applicable. Review source vouchers before finalizing base totals.",
                    base_qa_headers, base_qa_rows_values, formats, freeze_cols=6, tab_color="#C00000")

    open_gstr_headers = [
        "Source Sheet", "Source Row", "Event", "Transaction", "Supplier GSTIN", "Supplier", "Document No.", "Document Type",
        "Note Type", "Document Date", "2B Period", "Supplier Filing Date", "ITC Availability", "RCM", "Bucket",
        "Raw Base", "Signed Base Effect", "Signed IGST Effect", "Signed CGST Effect", "Signed SGST Effect", "Signed Cess Effect",
        "Signed GST Effect", "Linked Original ID", "Review Note",
    ]
    open_gstr_rows_values = [to_sheet_row(r, [
        "source_sheet", "source_row", "event", "transaction", "supplier_gstin", "supplier", "document_no", "document_type", "note_type",
        "document_date", "2b_period", "supplier_filing_date", "itc_availability", "rcm", "bucket", "raw_base", "signed_base_effect",
        "signed_igst_effect", "signed_cgst_effect", "signed_sgst_effect", "signed_cess_effect", "signed_gst_effect", "linked_original_id", "review_note",
    ]) for r in open_gstr]
    add_table_sheet(workbook, "Open 2B Items", "GSTR-2B records not located in books / separately ineligible",
                    "Contains unmatched 2B records, IMS-rejected items, and amendment events without a linked book original. Available, RCM, not-available, and rejected buckets are kept separate.",
                    open_gstr_headers, open_gstr_rows_values, formats, freeze_cols=5, tab_color="#BF9000")

    flow_headers = [
        "Transaction", "2B / ITC Bucket", "Book Month", "2B Period", "Event", "Timing",
        "Document Count", "Book Signed Base", "Book Signed GST", "2B Signed Base / Adjustment",
        "2B Signed IGST", "2B Signed CGST", "2B Signed SGST", "2B Signed Cess", "2B Signed GST",
        "Book less 2B GST",
    ]
    flow_rows_values = [to_sheet_row(r, [
        "transaction", "bucket", "book_month", "2b_period", "event", "timing", "count", "book_base", "book_gst",
        "b2b_base", "b2b_igst", "b2b_cgst", "b2b_sgst", "b2b_cess", "b2b_gst", "gst_variance",
    ]) for r in flow_rows]
    add_table_sheet(workbook, "Month Flow", "Matched documents flowing between book month and 2B period",
                    "Original matched documents show both book and 2B signed amounts. Amendments show only the 2B adjustment in the amendment period (book amount is not posted twice). Use this to explain timing differences across months.",
                    flow_headers, flow_rows_values, formats, freeze_cols=4, tab_color="#A5A5A5")

    control_headers = [
        "Control", "Raw detail source", "Raw row count", "Raw base", "Raw IGST", "Raw CGST", "Raw SGST", "Raw Cess",
        "2B summary IGST", "2B summary CGST", "2B summary SGST", "2B summary Cess", "IGST difference", "CGST difference",
        "SGST difference", "Cess difference", "Control Status", "Comment / treatment",
    ]
    control_rows_values = [to_sheet_row(r, [
        "control", "detail_source", "raw_count", "raw_base", "raw_igst", "raw_cgst", "raw_sgst", "raw_cess",
        "summary_igst", "summary_cgst", "summary_sgst", "summary_cess", "diff_igst", "diff_cgst", "diff_sgst", "diff_cess",
        "status", "comment",
    ]) for r in control_rows]
    add_table_sheet(workbook, "Source Control", "Source-detail vs GSTR-2B summary controls",
                    "Controls compare raw transaction-sheet tax amounts to the corresponding quarterly GSTR-2B summary section (before applying note signs). REVIEW rows are source-data/category differences that need user verification.",
                    control_headers, control_rows_values, formats, freeze_cols=2, tab_color="#7030A0")

    # Method & notes sheet: keep logic and known caveats visible inside the file.
    method_rows = [
        ("Period / entity", "FY 2025-26. GSTR-2B recipient GSTIN shown in the report: 24AAJCD4457C1ZV. Source report generation date: 29-Sep-2026."),
        ("Book month", "The Date column in each Purchase Register and Debit Note Register is used as the posting month."),
        ("2B month", "The supplier GSTR-1/IFF/1A/GSTR-5 Period field on the 2B transaction is used as the 2B period. Supplier filing date is shown separately."),
        ("Invoice matching", "Exact GSTIN + normalized Voucher No. ↔ 2B invoice number is tried first; Supplier Invoice No. is a fallback. Only after exact-GSTIN matching, a cross-GSTIN fallback may link an otherwise-unmatched invoice when the normalized document reference and calculated book taxable base / 2B taxable base (rounded to paise) occur exactly once on each complete source side. The cross-GSTIN match is explicitly flagged for review. Any duplicate/ambiguous reference-and-base key is left unmatched. No name-only fuzzy matching."),
        ("Credit-note matching", "Exact GSTIN + normalized book debit-note Voucher No. ↔ 2B note number is tried first; Voucher Ref No. is a fallback. The same unique reference + paise-exact taxable-base cross-GSTIN rule applies only between book debit notes and 2B supplier notes, and is flagged for review. Duplicate/ambiguous keys remain unmatched."),
        ("Normalization", "Reference normalization converts to uppercase and removes spaces/punctuation only. GSTIN is retained as part of every key."),
        ("Book base", "Calculated book base = Gross Total + source All Items − CGST − SGST − IGST − Cess − Round Off. For Purchase Register rows only, Purchase Accounts is added back when the ledger-column sum exceeds this bridge by approximately twice the source Purchase Accounts amount (₹0.05 tolerance), the observed contra/withholding pattern. This is a targeted source-data rule, not a general posting assumption. The raw ledger-column sum and ledger-sum-minus-base difference remain on Book Invoice Recon / Book Credit Note Recon / Open Book Items, with differences over ₹1 listed on Book Base QA. Review that tab before treating totals as final."),
        ("Book Base QA", "Rows on Book Base QA have a ledger-column-sum difference greater than ₹1 from the calculated base. Rows where the Purchase Accounts pattern was applied are marked; verify those ledger postings against source vouchers. Other rows are review exceptions, not automatically errors."),
        ("No separate GST", "Where book CGST+SGST+IGST/Cess columns are zero, the book ledger amount is shown in the no-separate-GST base. It may be non-GST, exempt/nil-rated, ineligible/embedded tax, or tax not separately posted; the source files do not distinguish these reasons."),
        ("Book debit notes", "The Debit Note Register is signed negative for purchase/ITC reconciliation, based on the uploaded samples matching supplier credit notes in GSTR-2B. Please verify unmatched debit-note entries; if any are buyer-issued debit notes, their sign should be reviewed."),
        ("2B note sign", "Supplier credit note type C is signed negative (reduces input credit). Supplier debit note type D is signed positive. Raw positive values remain visible in the GSTR2B Detail sheet."),
        ("Amendments", "B2BA and B2B-CDNRA are normalized as revised value less the matching original value and placed in the amended supplier period. This avoids double counting the full revised document after the original was already reflected. If the original is absent, the revised value is treated as a full addition and flagged."),
        ("Availability buckets", "ITC Availability=Yes and non-RCM is compared as standard available 2B ITC. RCM, ITC Availability=No, and IMS-rejected records are reported separately and not mixed into standard available ITC."),
        ("RCM", "The GSTR-2B report's RCM advisory says the tax is for reverse-charge reporting and ITC may be availed after payment. RCM effect is shown separately; it is not mixed into the primary standard-available variance."),
        ("Monthly difference", "Book less 2B = book net CGST/SGST/IGST ledger postings (purchases plus signed debit notes) minus standard non-RCM available GSTR-2B signed effect. Positive means book higher; negative means 2B higher. The base variance excludes book entries with no separate GST ledger posting."),
        ("Timing", "Month Flow lists book posting month vs supplier 2B period for each matched group. This exposes prior-month book items appearing in a later 2B period and items whose 2B period predates the book posting."),
        ("Open items", "Open Book Items / Open 2B Items are follow-up lists. An unmatched item can mean a period delay, differing reference, non-GST purchase, unavailable credit, or missing record; it is not automatically an error."),
        ("Cess", "Book registers expose CGST/SGST/IGST only. GSTR-2B Cess is included separately in detail and in the total 2B tax effect; any non-zero cess would require review."),
        ("Review warning", "Source Control deliberately flags raw 2B note records whose availability/RCM status does not tie to the quarterly ITC summary. Verify these with the downloaded portal statement/IMS before filing or claiming ITC."),
        ("Tax use", "This workbook is a data reconciliation aid, not a tax/legal opinion or confirmation of ITC eligibility. Validate classifications, RCM payment, blocked/ineligible credit, and return treatment with your tax professional."),
        ("Inputs", ", ".join(PURCHASE_FILES + DEBIT_NOTE_FILES + [GSTR_FILE])),
    ]
    method_headers = ["Topic", "Logic / note"]
    add_table_sheet(workbook, "Method & Notes", "Methodology, assumptions and review cautions",
                    "Read this sheet before relying on the monthly variance. All source row/file references are retained in the detail tabs.",
                    method_headers, method_rows, formats, freeze_cols=1, widths=[27, 120], tab_color="#17365D")

    # Keep the dashboard selected when the workbook opens.
    dashboard_ws.activate()
    dashboard_ws.set_first_sheet()
    workbook.close()


def main():
    for filename in PURCHASE_FILES + DEBIT_NOTE_FILES + [GSTR_FILE]:
        if not (ROOT / filename).exists():
            raise FileNotFoundError(f"Missing source workbook: {filename}")
    purchases = []
    book_notes = []
    for f in PURCHASE_FILES:
        purchases.extend(read_book_register(f, "Purchase"))
    for f in DEBIT_NOTE_FILES:
        book_notes.extend(read_book_register(f, "Debit Note"))

    wb = read_gstr_workbook()
    b2b, notes, raw_amendments, rejected = extract_gstr_records(wb)

    invoice_matches, invoice_reverse, _ = match_one_to_one(
        purchases, b2b, "voucher", ["invoice"], uniqueness_gstr_records=b2b + rejected)
    note_matches, note_reverse, _ = match_one_to_one(book_notes, notes, "voucher", ["invoice"])
    # Put the match fields on GSTR original records for later amendment links.
    for g in b2b:
        info = invoice_matches.get(g["id"])
        if info:
            book = info["book"]
            g["book_id"], g["book_record"] = book["id"], book
            g["match_method"] = info["method"]
            g["match_duplicate"] = info["duplicate"]
            g["match_ambiguous"] = info["ambiguous"]
            g["match_gstin_mismatch"] = info.get("gstin_mismatch", False)
        else:
            g["book_id"], g["book_record"] = "", None
    for g in notes:
        info = note_matches.get(g["id"])
        if info:
            book = info["book"]
            g["book_id"], g["book_record"] = book["id"], book
            g["match_method"] = info["method"]
            g["match_duplicate"] = info["duplicate"]
            g["match_ambiguous"] = info["ambiguous"]
            g["match_gstin_mismatch"] = info.get("gstin_mismatch", False)
        else:
            g["book_id"], g["book_record"] = "", None
    for b in purchases:
        if b["id"] in invoice_reverse:
            b["match_info"] = {"gstr": invoice_reverse[b["id"]]}
    for b in book_notes:
        if b["id"] in note_reverse:
            b["match_info"] = {"gstr": note_reverse[b["id"]]}

    rejected_pool = [dict(b) for b in purchases if not b.get("match_info")]
    rejected_matches, rejected_reverse, _ = match_one_to_one(
        rejected_pool, rejected, "voucher", ["invoice"],
        uniqueness_book_records=purchases, uniqueness_gstr_records=b2b + rejected)
    purchase_by_id = {b["id"]: b for b in purchases}
    for g in rejected:
        info = rejected_matches.get(g["id"])
        if info:
            matched_copy = info["book"]
            b = purchase_by_id[matched_copy["id"]]
            g["book_id"], g["book_record"] = b["id"], b
            g["match_method"] = info["method"]
            g["match_duplicate"] = info["duplicate"]
            g["match_ambiguous"] = info["ambiguous"]
            b["rejected_match"] = g
        else:
            g["book_id"], g["book_record"] = "", None
    # Replace the copy-based reverse map with the original purchase objects' IDs.
    rejected_reverse = {bid: g for bid, g in rejected_reverse.items()}
    for a in raw_amendments:
        a["book_id"] = ""
        a["book_record"] = None

    amendments = normalise_amendments(raw_amendments, b2b, notes)
    attach_amendment_book_links(amendments, invoice_matches, purchases, book_notes)

    # Associate each amendment with its original record for the book-level detail.
    amendments_by_original = defaultdict(list)
    for a in amendments:
        if a.get("linked_original_id"):
            amendments_by_original[a["linked_original_id"]].append(a)

    # Rejected matches are separate from available 2B matches; `match_one_to_one`
    # has already set duplicate flags on its pool. Preserve those flags.
    for b in purchases + book_notes:
        b.setdefault("amendments", [])
        if not b.get("match_info"):
            b["match_info"] = None

    # Build details. Convert the GSTR original event references used by matching
    # to explicit `book_id` links and leave the source objects otherwise intact.
    for g in b2b + notes:
        if g.get("book_record"):
            g["book_record"]["match_info"]["gstr"] = g
    for a in amendments:
        if a.get("original") and a["original"].get("book_record"):
            a["book_record"] = a["original"]["book_record"]
            a["book_id"] = a["book_record"]["id"]
            a["book_month"] = a["book_record"]["month"]
    for e in rejected:
        if e.get("book_record"):
            e["book_record"].setdefault("rejected_match", e)

    months = build_monthly_summary(purchases, book_notes, b2b, notes, amendments, rejected)
    flow = build_month_flow(b2b, notes, amendments)
    control_rows = build_source_controls(wb, b2b, notes, raw_amendments, rejected)

    invoice_rows = build_recon_rows(purchases, invoice_reverse, rejected_reverse, {r["id"]: r for r in b2b}, amendments_by_original)
    note_rows = build_credit_note_rows(book_notes, note_reverse, amendments_by_original)
    gstr_rows = build_gstr_detail_rows(b2b, notes, amendments, rejected)
    open_book = build_open_book_rows(purchases, book_notes)
    base_qa = build_book_base_qa(purchases, book_notes)
    open_gstr = build_open_gstr_rows(b2b, notes, amendments, rejected)

    # Counts are one-to-one pairs, not raw key overlaps.
    counts = {
        "purchase_rows": len(purchases), "note_rows": len(book_notes), "b2b_rows": len(b2b),
        "note_2b_rows": len(notes), "invoice_amendments": sum(1 for a in amendments if a["source_sheet"] == "B2BA"),
        "note_amendments": sum(1 for a in amendments if a["source_sheet"] == "B2B-CDNRA"),
        "rejected_rows": len(rejected),
        "matched_invoice_pairs": sum(1 for g in b2b if g.get("book_id")),
        "matched_note_pairs": sum(1 for g in notes if g.get("book_id")),
        "gstin_mismatch_matches": sum(1 for g in b2b + notes if g.get("match_gstin_mismatch")),
        "open_book_rows": len(open_book), "base_qa_rows": len(base_qa),
        "open_2b_original_rows": sum(1 for g in b2b + notes if not g.get("book_id")),
    }
    summary = {
        "book_gst": sum(r["book_net_gst"] for r in months),
        "eligible_gst": sum(r["b2b_available_gst"] for r in months),
        "difference_gst": sum(r["diff_gst"] for r in months),
        "no_tax_base": sum(r["book_no_tax_base_net"] for r in months),
        "rcm_gst": sum(r["b2b_rcm_gst"] for r in months),
    }

    # Source summary control requires the raw values before amendment
    # normalization; add bucket detail and source status checks above.
    write_workbook(months, invoice_rows, note_rows, gstr_rows, open_book, base_qa, open_gstr, flow, control_rows, counts, summary)
    wb.release_resources()

    print(f"Created {OUTPUT}")
    print(f"Book purchase rows: {len(purchases):,}; book debit notes: {len(book_notes):,}")
    print(f"GSTR-2B B2B: {len(b2b):,}; supplier notes: {len(notes):,}; amendments: {len(amendments):,}; rejected: {len(rejected):,}")
    print(f"Matched invoice pairs: {counts['matched_invoice_pairs']:,}; matched note pairs: {counts['matched_note_pairs']:,}")
    print(f"Open book rows: {len(open_book):,}; book base QA rows (>₹1 ledger-sum difference): {len(base_qa):,}; open 2B originals: {counts['open_2b_original_rows']:,}")
    print(f"Net book GST: {summary['book_gst']:,.2f}; 2B standard available signed GST: {summary['eligible_gst']:,.2f}; difference: {summary['difference_gst']:,.2f}")
    print(f"No-separate-GST book base (net): {summary['no_tax_base']:,.2f}; RCM GST effect (separate): {summary['rcm_gst']:,.2f}")
    print("Source controls:")
    for row in control_rows:
        if row["status"] != "OK":
            print(f"  REVIEW: {row['control']} | Δ IGST {row['diff_igst']:,.2f}, CGST {row['diff_cgst']:,.2f}, SGST {row['diff_sgst']:,.2f}")


if __name__ == "__main__":
    main()
